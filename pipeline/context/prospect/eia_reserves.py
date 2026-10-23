"""Oil and gas: EIA proved reserves by state, year-end 2024 (a small table, joined to state polygons).

Source: "U.S. Crude Oil and Natural Gas Proved Reserves, Year-end 2024" (Form EIA-23L), released
2026-03-25: ``ARR_2024_TABLES_ALL.xlsx`` on https://www.eia.gov/naturalgas/crudeoilreserves/.

- sheet "6": crude oil and lease condensate proved reserves by state and area, million barrels;
- sheet "8": natural gas, wet after lease separation, proved reserves by state and area, billion cubic feet.

Each sheet lists states, their subdivisions indented beneath them ("   RRC District 8"), federal offshore
areas, "Lower 48 States", "Other States" and "U.S. Total". The columns used: total proved reserves at
12/31/2023 and at 12/31/2024, proved producing and proved nonproducing at 12/31/2024, and estimated
production in 2024. "-" means no data reported and "W" withheld; both become null with the reason kept.

The two sheets do not list the same areas. In the year-end 2024 file Pennsylvania (95,135 Bcf, about 16%
of US proved gas) and New York appear on the gas sheet only, and Nebraska and South Dakota on the crude
sheet only (each is folded into "Other States" on the other sheet), so the table is the union of both
sheets' rows; the side a row is missing from is null with a note saying so.

`row_kind` says what a row is: `state` (one of the 50 states or DC), `subdivision` (indented under a
state), `federal_offshore` (the "Federal Offshore" row and its areas), `other_states` (the residual row)
or `total` ("Lower 48 States", "U.S. Total"). Only `state` rows join to a polygon.

Outputs: ``oil_gas_reserves_by_state.csv`` (every row, subdivisions and offshore areas included: the
table), and the ``oil_gas_reserves_state`` GeoParquet layer -- one row per state with reserves, carrying
the vendored Census state polygon (public domain, `us.census.cartographic_boundaries`, whose provenance
the layer names beside EIA's) so the sampler answers "this state's reserves" at a point. A statewide
total, never a value for a site.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import io
import json
import pathlib
import re
import time
from dataclasses import dataclass
from typing import Any

import pyarrow as pa  # type: ignore[import-untyped, unused-ignore]

from pipeline.connectors.base import ParseError
from pipeline.connectors.http import HttpFailed, PoliteSession
from pipeline.connectors.store import Store
from pipeline.context.prospect.common import (
    REGIONS,
    FetchLog,
    now_iso,
    output_dir,
    provenance_for,
    snapshot_time,
    verify_terms,
)
from pipeline.context.prospect.coverage import CENSUS_SOURCE_ID, CENSUS_STATE_URL, public, region_box
from pipeline.context.prospect.eia_oil_gas import LICENCE_NAME, TERMS_PHRASES, TERMS_URL
from pipeline.context.prospect.geoparquet import write_geoparquet

SOURCE_ID = "us.eia.proved_reserves"
REPORT_YEAR = 2024
URL = f"https://www.eia.gov/naturalgas/crudeoilreserves/excel/ARR_{REPORT_YEAR}_TABLES_ALL.xlsx"
VINTAGE = f"Proved reserves at year-end {REPORT_YEAR} (EIA-23L), released 2026-03-25"
ROOT = pathlib.Path(__file__).resolve().parents[3]
STATES_GEOJSON = ROOT / "data" / "vendored" / "regions" / "us_states.geojson"
#: Where the vendored state polygons came from (pipeline/context/regions.py STATE_URL), named in the
#: layer's provenance beside EIA's because the geometry is Census's, not EIA's.
GEOMETRY_PROVENANCE = {
    "geometry_source_id": CENSUS_SOURCE_ID,
    "geometry_source_url": CENSUS_STATE_URL,
    "geometry_licence": "Public domain (US Census Bureau, 17 U.S.C. §105)",
    "geometry_vintage": "2024 cartographic boundary file, 1:20,000,000 (data/vendored/regions)",
}
CRUDE_SHEET, GAS_SHEET = "6", "8"
#: The 50 states and DC as EIA's tables spell them: only these rows are `row_kind: state`.
STATE_NAMES = frozenset(
    {
        "Alabama", "Alaska", "Arizona", "Arkansas", "California", "Colorado", "Connecticut", "Delaware",
        "District of Columbia", "Florida", "Georgia", "Hawaii", "Idaho", "Illinois", "Indiana", "Iowa",
        "Kansas", "Kentucky", "Louisiana", "Maine", "Maryland", "Massachusetts", "Michigan", "Minnesota",
        "Mississippi", "Missouri", "Montana", "Nebraska", "Nevada", "New Hampshire", "New Jersey",
        "New Mexico", "New York", "North Carolina", "North Dakota", "Ohio", "Oklahoma", "Oregon",
        "Pennsylvania", "Rhode Island", "South Carolina", "South Dakota", "Tennessee", "Texas", "Utah",
        "Vermont", "Virginia", "Washington", "West Virginia", "Wisconsin", "Wyoming",
    }
)  # fmt: skip
FEDERAL_OFFSHORE = "Federal Offshore"
TOTAL_ROWS = frozenset({"Lower 48 States", "U.S. Total"})
#: Column positions in sheets 6 and 8 (row 5 headers, row 6 dates).
COLUMNS = {
    "proved_prev": 1,
    "production": 6,
    "proved_producing": 7,
    "proved_nonproducing": 8,
    "proved": 9,
}


@dataclass
class ReserveRow:
    area: str
    parent: str | None
    #: Not indented on the sheet: a state, or an aggregate ("Lower 48 States", "Federal Offshore", ...).
    top_level: bool
    values: dict[str, float | None]
    notes: dict[str, str]


def row_kind(area: str, parent: str | None) -> str:
    """`state`, `subdivision`, `federal_offshore`, `other_states` or `total` (module docstring)."""
    if parent is not None:
        return "federal_offshore" if parent == FEDERAL_OFFSHORE else "subdivision"
    if area in STATE_NAMES:
        return "state"
    if area == FEDERAL_OFFSHORE:
        return "federal_offshore"
    if area in TOTAL_ROWS:
        return "total"
    if area.lower().startswith("other states"):
        return "other_states"
    return "other"


def _cell(value: Any) -> tuple[float | None, str | None]:
    if value is None:
        return None, None
    if isinstance(value, int | float):
        return float(value), None
    s = str(value).strip()
    if s in ("-", "--"):
        return None, "no data reported"
    if s.upper() == "W":
        return None, "withheld to avoid disclosure"
    if s.upper() == "NA":
        return None, "not available"
    try:
        return float(s.replace(",", "")), None
    except ValueError:
        return None, f"unparsed: {s}"


def parse_sheet(rows: list[tuple[Any, ...]]) -> list[ReserveRow]:
    """Rows of sheet 6 or 8 -> one ReserveRow per state, subdivision or area (notes and totals kept)."""
    header = next(
        (i for i, r in enumerate(rows) if r and str(r[0] or "").strip().lower().startswith("state and")), None
    )
    if header is None:
        raise ParseError("no 'State and Subdivision' header row")
    out: list[ReserveRow] = []
    parent: str | None = None
    for r in rows[header + 1 :]:
        if not r or r[0] is None:
            continue
        label = str(r[0])
        name = label.strip()
        if not name or re.match(r"^\(?[a-z]\)|^notes?:|^source:|^-\s", name, re.I):
            continue
        if len(r) <= COLUMNS["proved"]:
            continue
        indented = label.startswith(" ")
        values: dict[str, float | None] = {}
        notes: dict[str, str] = {}
        for key, idx in COLUMNS.items():
            v, note = _cell(r[idx])
            values[key] = v
            if note:
                notes[key] = note
        if all(v is None for v in values.values()) and not notes:
            continue
        name = re.sub(r"\s*\([a-z]\)$", "", name)
        if indented:
            out.append(ReserveRow(name, parent, False, values, notes))
        else:
            parent = name
            out.append(ReserveRow(name, None, True, values, notes))
    if not any(r.area == "Texas" for r in out):
        raise ParseError("no Texas row: layout changed")
    return out


def parse(content: bytes) -> dict[str, list[ReserveRow]]:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several unrelated types for a non-xlsx body
        raise ParseError(f"not an xlsx workbook: {exc}") from exc
    names = wb.sheetnames
    if CRUDE_SHEET not in names or GAS_SHEET not in names:
        raise ParseError(f"sheets {CRUDE_SHEET!r}/{GAS_SHEET!r} missing: {names}")
    return {
        "crude": parse_sheet([tuple(r) for r in wb[CRUDE_SHEET].iter_rows(values_only=True)]),
        "gas": parse_sheet([tuple(r) for r in wb[GAS_SHEET].iter_rows(values_only=True)]),
    }


Key = tuple[str, str | None]


def _merged_keys(crude: list[ReserveRow], gas: list[ReserveRow]) -> list[Key]:
    """Every (area, parent) on either sheet, once: the crude sheet's order, with each row found only on
    the gas sheet placed after the row that precedes it there (Pennsylvania after Oklahoma)."""
    keys: list[Key] = [(r.area, r.parent) for r in crude]
    seen = set(keys)
    previous: Key | None = None
    for r in gas:
        key = (r.area, r.parent)
        if key not in seen:
            at = keys.index(previous) + 1 if previous is not None else 0
            # A new top-level row goes after the previous row's subdivisions, not between them.
            if r.parent is None:
                while at < len(keys) and keys[at][1] is not None:
                    at += 1
            keys.insert(at, key)
            seen.add(key)
        previous = key
    return keys


def table(parsed: dict[str, list[ReserveRow]]) -> list[dict[str, Any]]:
    """Crude and gas rows joined on (area, parent) over the union of both sheets' rows: an area on one
    sheet only (Pennsylvania and New York are on the gas sheet only, year-end 2024) keeps its row, with
    the other side null and a note that says so."""
    crude = {(r.area, r.parent): r for r in parsed["crude"]}
    gas = {(r.area, r.parent): r for r in parsed["gas"]}
    out = []
    for key in _merged_keys(parsed["crude"], parsed["gas"]):
        c, g = crude.get(key), gas.get(key)
        area, parent = key
        kind = row_kind(area, parent)
        notes = [f"crude {k}: {v}" for k, v in (c.notes.items() if c else ())]
        notes += [f"gas {k}: {v}" for k, v in (g.notes.items() if g else ())]
        if c is None:
            notes.append("no row on the crude oil sheet (Table 6)")
        if g is None:
            notes.append("no row on the natural gas sheet (Table 8)")
        cv: dict[str, float | None] = c.values if c else {}
        gv: dict[str, float | None] = g.values if g else {}
        out.append(
            {
                "area": area,
                "parent_state": parent,
                "row_kind": kind,
                "is_state_row": kind == "state",
                "crude_oil_lease_condensate_mmbbl": cv.get("proved"),
                "crude_oil_lease_condensate_mmbbl_prev": cv.get("proved_prev"),
                "crude_proved_producing_mmbbl": cv.get("proved_producing"),
                "crude_proved_nonproducing_mmbbl": cv.get("proved_nonproducing"),
                "crude_production_mmbbl": cv.get("production"),
                "natural_gas_wet_bcf": gv.get("proved"),
                "natural_gas_wet_bcf_prev": gv.get("proved_prev"),
                "gas_proved_producing_bcf": gv.get("proved_producing"),
                "gas_proved_nonproducing_bcf": gv.get("proved_nonproducing"),
                "gas_production_bcf": gv.get("production"),
                "notes": "; ".join(notes),
                "report_year": REPORT_YEAR,
            }
        )
    return out


SCHEMA = [
    ("state", pa.string()),
    ("region_id", pa.string()),
    ("crude_oil_lease_condensate_mmbbl", pa.float64()),
    ("natural_gas_wet_bcf", pa.float64()),
    ("crude_proved_nonproducing_mmbbl", pa.float64()),
    ("gas_proved_nonproducing_bcf", pa.float64()),
    ("report_year", pa.int32()),
    ("notes", pa.string()),
]


def state_layer_rows(
    rows: list[dict[str, Any]], states: dict[str, Any], *, bbox: tuple[float, float, float, float]
) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """State rows that have a polygon and touch the region box -> (rows, geometries)."""
    out_rows, geoms = [], []
    for row in rows:
        if not row["is_state_row"]:
            continue
        feature = states.get(row["area"])
        if feature is None:
            continue
        gb = _geom_bbox(feature["geometry"])
        if gb[0] > bbox[2] or gb[2] < bbox[0] or gb[1] > bbox[3] or gb[3] < bbox[1]:
            continue
        out_rows.append(
            {
                "state": row["area"],
                "region_id": feature["properties"].get("region_id"),
                "crude_oil_lease_condensate_mmbbl": row["crude_oil_lease_condensate_mmbbl"],
                "natural_gas_wet_bcf": row["natural_gas_wet_bcf"],
                "crude_proved_nonproducing_mmbbl": row["crude_proved_nonproducing_mmbbl"],
                "gas_proved_nonproducing_bcf": row["gas_proved_nonproducing_bcf"],
                "report_year": row["report_year"],
                "notes": row["notes"] or None,
            }
        )
        geoms.append(feature["geometry"])
    return out_rows, geoms


def _geom_bbox(geom: dict[str, Any]) -> tuple[float, float, float, float]:
    polys = geom["coordinates"] if geom["type"] == "MultiPolygon" else [geom["coordinates"]]
    xs = [p[0] for poly in polys for p in poly[0]]
    ys = [p[1] for poly in polys for p in poly[0]]
    return min(xs), min(ys), max(xs), max(ys)


def build_from_bytes(
    content: bytes,
    *,
    region: str,
    out_dir: pathlib.Path,
    retrieved_at: str,
    states_path: pathlib.Path = STATES_GEOJSON,
) -> list[dict[str, Any]]:
    rows = table(parse(content))
    prov = provenance_for(
        SOURCE_ID, source_url=URL, retrieved_at=retrieved_at, licence=LICENCE_NAME, vintage=VINTAGE
    )
    csv_path = out_dir / "oil_gas_reserves_by_state.csv"
    with csv_path.open("w", newline="", encoding="utf-8") as fh:
        fields = list(rows[0]) + list(prov.columns())
        writer = csv.DictWriter(fh, fieldnames=fields)
        writer.writeheader()
        for row in rows:
            writer.writerow({**row, **prov.columns()})
    states = {f["properties"]["name"]: f for f in json.loads(states_path.read_text())["features"]}
    layer_rows, geoms = state_layer_rows(rows, states, bbox=REGIONS[region])
    out = out_dir / "oil_gas_reserves_state.parquet"
    coverage = region_box(region)
    summary = write_geoparquet(
        out,
        layer_rows,
        geoms,
        schema=SCHEMA,
        provenance=prov,
        extra_metadata={**GEOMETRY_PROVENANCE, "coverage": coverage},
    )
    summary["table_rows"] = len(rows)
    summary["table_file"] = csv_path.name
    summary["table_bytes"] = csv_path.stat().st_size
    return [
        {
            "layer_id": "oil_gas_reserves_state",
            "file": out.name,
            "kind": "polygons",
            "provenance": {**prov.record(), **GEOMETRY_PROVENANCE},
            "licence_name": LICENCE_NAME,
            "summary": summary,
            "extra_files": [csv_path.name],
            "coverage": public(coverage),
        }
    ]


def build(
    *,
    region: str = "texas",
    out_dir: pathlib.Path | None = None,
    snapshot: pathlib.Path | None = None,
    session: PoliteSession | None = None,
    store: Store | None = None,
) -> list[dict[str, Any]]:
    out_dir = output_dir(out_dir)
    if snapshot is not None:
        return build_from_bytes(
            snapshot.read_bytes(), region=region, out_dir=out_dir, retrieved_at=snapshot_time(snapshot)
        )
    session = session or PoliteSession(rate_limits={"www.eia.gov": 1.0})
    log = FetchLog(SOURCE_ID)
    verify_terms(session, TERMS_URL, TERMS_PHRASES, log=log)
    resp = session.get(URL)
    if resp.status_code != 200:
        raise HttpFailed(f"{URL} answered HTTP {resp.status_code}")
    retrieved_at = now_iso()
    log.add(
        url=URL,
        status=200,
        bytes=len(resp.content),
        last_modified=resp.headers.get("Last-Modified"),
        retrieved_at=retrieved_at,
    )
    if store is not None:
        from pipeline.connectors.store import ts_token

        store.write_snapshot(SOURCE_ID, ts_token(dt.datetime.now(dt.UTC)), "xlsx", resp.content)
    outputs = build_from_bytes(resp.content, region=region, out_dir=out_dir, retrieved_at=retrieved_at)
    log.write(store, region=region)
    return outputs


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    parser.add_argument("--region", choices=sorted(REGIONS), default="texas")
    parser.add_argument("--snapshot", type=pathlib.Path)
    parser.add_argument("--out-dir", type=pathlib.Path)
    args = parser.parse_args(argv)
    t0 = time.monotonic()
    outputs = build(region=args.region, out_dir=args.out_dir, snapshot=args.snapshot, store=Store())
    print(json.dumps({"layers": outputs, "elapsed_s": round(time.monotonic() - t0, 2)}))  # noqa: T201


if __name__ == "__main__":
    main()
