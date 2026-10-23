"""The price runner end to end against a fake network (`conftest.FakeHttp`) and a scratch store:
the gate, coverage-driven plans, conditional requests, snapshot re-use and ERCOT's
three-downloads rule, corrections, DQ holds and the run record (docs/29 §7)."""

from __future__ import annotations

import datetime as dt

import pandas as pd
import pytest

from pipeline.connectors.prices import ercot, ercot_api
from pipeline.connectors.prices import store as pstore
from pipeline.connectors.prices.base import Coverage, FetchRequest, PriceGateRefused, gate
from pipeline.connectors.prices.conftest import FakeHttp, fixture_bytes, snap
from pipeline.connectors.prices.ercot import parse_display_page, report_downloads
from pipeline.connectors.prices.runner import can_fork, run
from pipeline.connectors.store import Store

ERCOT = "us.iso.ercot.prices"
ELEXON = "gb.elexon.bmrs_prices"
EIA = "us.eia.retail_sales"
SPP = "us.iso.spp.prices"
NOW = dt.datetime(2026, 10, 22, 22, 0, tzinfo=dt.UTC)


@pytest.fixture
def store(tmp_path) -> Store:
    return Store(tmp_path)


def page(name: str, **headers: str):
    def answer(**_kw):
        return 200, fixture_bytes(name), dict(headers)

    return answer


# ------------------------------------------------------------------------------------------- gate
@pytest.mark.parametrize(
    ("source_id", "reason"),
    [
        ("us.iso.pjm.prices", "pending"),
        ("us.iso.miso.prices", "pending"),
        ("us.iso.isone.prices", "pending"),
        ("eu.entsoe.day_ahead_prices", "excluded"),
        ("gb.elexon.market_index_data", "pending"),
    ],
)
def test_the_gate_refuses_rows_that_could_never_be_published(source_id, reason, registry) -> None:
    with pytest.raises(PriceGateRefused, match=reason):
        gate(registry.get(source_id), frozenset({"open", "attribution", "noncommercial"}))


def test_spp_is_fetched_only_under_the_noncommercial_posture(registry, store) -> None:
    entry = registry.get(SPP)
    with pytest.raises(PriceGateRefused, match="noncommercial"):
        gate(entry, frozenset({"open", "attribution"}))
    gate(entry, frozenset({"open", "attribution", "noncommercial"}))
    http = FakeHttp({})
    with pytest.raises(PriceGateRefused):
        run(SPP, "daily", store=store, http=http, now=NOW, publishable=frozenset({"open", "attribution"}))
    assert http.calls == [] and store.runs(SPP) == []


# --------------------------------------------------------------------------------- daily coverage
def test_a_daily_run_asks_only_for_days_not_held(store) -> None:
    http = FakeHttp({"system-prices/2026-10-21": "elexon_system_prices_20261021.json"})
    first = run(ELEXON, "daily", store=store, http=http, now=NOW)
    assert first.status == "ok"
    assert first.record["rows_new"] == 48 and first.record["dq_status"] == "warn"
    asked = [c["url"].rsplit("/", 1)[-1] for c in http.calls]
    assert asked[0] == "2026-10-08" and asked[-1] == "2026-10-21" and len(asked) == 14
    warn = {c["check"]: c for c in first.record["dq"]["checks"]}["request_failures"]
    assert warn["level"] == "warn" and len(warn["data"]["failures"]) == 13

    stored = pstore.read_month(store, ELEXON, "2026-10")
    assert stored is not None and len(stored) == 48
    snaps = list((store.root / "snapshots" / ELEXON).iterdir())
    assert len(snaps) == 1 and snaps[0].name.endswith(".rt-20261021.json")

    http2 = FakeHttp({"system-prices/2026-10-21": "elexon_system_prices_20261021.json"})
    second = run(ELEXON, "daily", store=store, http=http2, now=NOW + dt.timedelta(hours=1))
    assert not any("2026-10-21" in c["url"] for c in http2.calls)  # complete: not asked again
    assert len(http2.calls) == 13
    assert second.status == "failed"  # every request this time failed
    assert len(store.runs(ELEXON)) == 2


def test_the_run_record_is_the_scheduler_s_shape(store) -> None:
    http = FakeHttp({"system-prices/2026-10-21": "elexon_system_prices_20261021.json"})
    rec = run(ELEXON, "daily", store=store, http=http, now=NOW, trigger="schedule").record
    for key in (
        "id",
        "source_id",
        "mode",
        "trigger",
        "parser_version",
        "started_at",
        "finished_at",
        "status",
        "requests",
        "requests_made",
        "rows_seen",
        "rows_new",
        "rows_changed",
        "dq_status",
        "bytes",
    ):
        assert key in rec, key
    assert rec["kind"] == "price" and rec["trigger"] == "schedule" and rec["requests_made"] == 14
    with pytest.raises(ValueError):
        run(ELEXON, "weekly", store=store, http=http, now=NOW)
    with pytest.raises(ValueError):
        run("us.eia.retail_sales", "latest", store=store, http=http, now=NOW)


# ---------------------------------------------------------------------- conditional + corrections
def test_latest_uses_validators_and_a_changed_price_is_a_correction(store) -> None:
    url = "cdr/html/real_time_spp.html"
    body = fixture_bytes("ercot_real_time_spp.html")
    http = FakeHttp(
        {
            url: page(
                "ercot_real_time_spp.html", ETag='"abc"', **{"Last-Modified": "Thu, 22 Oct 2026 21:50:00 GMT"}
            )
        }
    )
    first = run(ERCOT, "latest", store=store, http=http, now=NOW)
    assert first.status == "ok" and first.record["rows_new"] > 0
    assert "If-None-Match" not in http.calls[0]["headers"]

    def not_modified(**kw):
        assert kw["headers"]["If-None-Match"] == '"abc"'
        assert kw["headers"]["If-Modified-Since"] == "Thu, 22 Oct 2026 21:50:00 GMT"
        return 304, b"", {}

    second = run(
        ERCOT, "latest", store=store, http=FakeHttp({url: not_modified}), now=NOW + dt.timedelta(minutes=15)
    )
    assert second.status == "unchanged" and second.record["rows_new"] == 0 and second.record["bytes"] == 0
    assert len(list((store.root / "snapshots" / ERCOT).iterdir())) == 1  # a 304 stores nothing

    header, rows, _ = parse_display_page(body)
    first_price = rows[0][2]
    corrected = body.replace(f">{first_price}<".encode(), f">{float(first_price) + 100:.2f}<".encode(), 1)
    third = run(
        ERCOT,
        "latest",
        store=store,
        http=FakeHttp({url: lambda **_k: (200, corrected, {"ETag": '"abd"'})}),
        now=NOW + dt.timedelta(minutes=30),
    )
    assert third.status == "ok" and third.record["rows_changed"] == 1 and third.record["rows_new"] == 0
    log = store.read_parquet(store.root / third.record["outputs"]["corrections"])
    assert len(log) == 1
    row = log.iloc[0]
    assert row["location"] == header[2] and row["price_after"] == pytest.approx(float(first_price) + 100)
    assert row["price_before"] == pytest.approx(float(first_price))
    stored = pstore.read_month(store, ERCOT, "2026-10")
    assert stored is not None and stored["record_id"].is_unique


# ------------------------------------------------------------------------------------------ DQ
def test_a_retail_table_that_shrinks_by_half_is_held(store) -> None:
    http = FakeHttp(
        {
            "sales_revenue.xlsx": page(
                "eia_sales_revenue_trimmed.xlsx", **{"Last-Modified": "Thu, 24 Sep 2026 13:48:22 GMT"}
            )
        }
    )
    first = run(EIA, "daily", store=store, http=http, now=NOW)
    assert first.status == "ok" and first.record["rows_new"] == 6 * 13 * 5

    def not_modified(**kw):
        assert kw["headers"]["If-Modified-Since"] == "Thu, 24 Sep 2026 13:48:22 GMT"
        return 304, b"", {}

    second = run(EIA, "daily", store=store, http=FakeHttp({"sales_revenue.xlsx": not_modified}), now=NOW)
    assert second.status == "unchanged" and second.record["dq_status"] == "pass"

    import io

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(fixture_bytes("eia_sales_revenue_trimmed.xlsx")))
    ws = wb["Monthly-States"]
    for row in range(40, ws.max_row + 1):  # keep the header and about half the state rows
        for col in range(1, ws.max_column + 1):
            ws.cell(row=row, column=col).value = None
    buf = io.BytesIO()
    wb.save(buf)
    shrunk = buf.getvalue()
    third = run(
        EIA,
        "daily",
        store=store,
        http=FakeHttp({"sales_revenue.xlsx": lambda **_k: (200, shrunk, {})}),
        now=NOW,
    )
    assert third.status == "partial" and third.record["hold_reasons"] == ["row_count"]
    table = store.read_parquet(pstore.retail_path(store, EIA))
    assert len(table) == 6 * 13 * 5  # the held run wrote nothing


def test_a_page_with_none_of_the_served_locations_is_held(store) -> None:
    body = fixture_bytes("ercot_real_time_spp.html")
    for name in (
        "HB_BUSAVG",
        "HB_HOUSTON",
        "HB_HUBAVG",
        "HB_NORTH",
        "HB_PAN",
        "HB_SOUTH",
        "HB_WEST",
        "LZ_AEN",
        "LZ_CPS",
        "LZ_HOUSTON",
        "LZ_LCRA",
        "LZ_NORTH",
        "LZ_RAYBN",
        "LZ_SOUTH",
        "LZ_WEST",
    ):
        body = body.replace(f">{name}<".encode(), f">X{name}<".encode())
    res = run(
        ERCOT,
        "latest",
        store=store,
        http=FakeHttp({"cdr/html/real_time_spp.html": lambda **_k: (200, body, {})}),
        now=NOW,
    )
    assert res.status == "partial" and "location_vocabulary" in res.record["hold_reasons"]
    assert pstore.list_months(store, ERCOT) == []
    assert len(list((store.root / "snapshots" / ERCOT).iterdir())) == 1  # the evidence is kept


# --------------------------------------------------------------------------- ERCOT annual files
ANNUAL = {
    "reportTypeId=13061": "ercot_doclist_13061.json",
    "reportTypeId=13060": "ercot_doclist_13060.json",
    "doclookupId=1177737535": "ercot_rtm_annual_2025_nov02.zip",
    "doclookupId=1287900508": "ercot_rtm_annual_2026_sample.zip",
    "doclookupId=1177667469": "ercot_dam_annual_2025_nov02.zip",
    "doclookupId=1287898588": "ercot_dam_annual_2026_sample.zip",
}


def test_ercot_backfill_downloads_each_annual_workbook_once(store, monkeypatch) -> None:
    # A 70-day window from 2026-01-03 reaches back to 2025-10-25, so the 2025 workbooks' November
    # sheet (2025-11-02 recorded) is inside it; the window spans 2025 and 2026.
    monkeypatch.setattr(ercot.Connector, "backfill_days", 70)
    now = dt.datetime(2026, 1, 3, 12, 0, tzinfo=dt.UTC)
    http = FakeHttp(ANNUAL)
    first = run(ERCOT, "backfill", store=store, http=http, now=now, trigger="backfill")
    downloads = [c for c in http.calls if "mirDownload" in c["url"]]
    assert len(downloads) == 4 and sum("IceDocListJsonWS" in c["url"] for c in http.calls) == 2
    annual = [e for e in first.record["requests"] if (e.get("meta") or {}).get("kind") == "annual"]
    assert {e["meta"]["report"] for e in annual} == {
        "RTMLZHBSPP_2025",
        "RTMLZHBSPP_2026",
        "DAMLZHBSPP_2025",
        "DAMLZHBSPP_2026",
    }
    assert all(e["downloaded"] and not e["reused"] and e["sha256"] for e in annual)
    assert first.status == "ok" and first.record["rows_new"] == 1500 + 2820 + 375 + 705
    ledger = pstore.read_ledger(store, ERCOT)
    assert sorted(e["report"] for e in ledger) == sorted(e["meta"]["report"] for e in annual)
    assert all(e["state"] == "stored" and e["sha256"] and e["object_key"] for e in ledger)
    assert pstore.backfill_attempts(store, ERCOT) == {first.run_path.stem}  # the marker's ts
    # The display pages for the window's days are asked after the workbooks (all 404 here).
    assert any("_real_time_spp.html" in c["url"] for c in http.calls)

    http2 = FakeHttp(ANNUAL)
    second = run(
        ERCOT, "backfill", store=store, http=http2, now=now + dt.timedelta(hours=1), trigger="backfill"
    )
    assert not any("mirDownload" in c["url"] for c in http2.calls)
    again = [e for e in second.record["requests"] if (e.get("meta") or {}).get("kind") == "annual"]
    assert len(again) == 4 and all(e["reused"] and not e["downloaded"] for e in again)
    assert second.status == "unchanged"

    # A stored workbook that fails its SHA-256 is downloaded again.
    target = next(e for e in annual if e["meta"]["report"] == "DAMLZHBSPP_2025")
    store.backend.put(target["object_key"], b"corrupt")
    http3 = FakeHttp(ANNUAL)
    run(ERCOT, "backfill", store=store, http=http3, now=now + dt.timedelta(hours=2), trigger="backfill")
    assert [c["url"].rsplit("=", 1)[-1] for c in http3.calls if "mirDownload" in c["url"]] == ["1177667469"]


def test_a_report_downloaded_three_times_in_twelve_months_is_not_asked_a_fourth(registry) -> None:
    now = dt.datetime(2026, 10, 22, tzinfo=dt.UTC)
    stamps = ["2026-01-05T00:00:00+00:00", "2026-05-05T00:00:00+00:00", "2026-09-05T00:00:00+00:00"]
    assert report_downloads({"RTMLZHBSPP_2026": stamps}, "RTMLZHBSPP_2026", now) == 3
    assert report_downloads({"RTMLZHBSPP_2026": ["2025-09-05T00:00:00+00:00"]}, "RTMLZHBSPP_2026", now) == 0

    raw = snap("ercot_doclist_13061.json")
    request = FetchRequest(key="list-13061", url="u", ext="json", meta={"kind": "doc_list", "market": "rt"})
    cov = Coverage(downloads={"RTMLZHBSPP_2026": stamps})

    plain = ercot.Connector(registry.get(ERCOT), env={}, clock=lambda: now)
    follow = plain.follow_up(request, raw, now.date(), cov)
    assert [r.meta["report"] for r in follow] == ["RTMLZHBSPP_2025"]
    assert plain.skipped == [
        {"report": "RTMLZHBSPP_2026", "reason": "downloaded 3 times in 12 months; no API account"}
    ]

    creds = dict.fromkeys(ercot.API_ENV, "x")
    keyed = ercot.Connector(registry.get(ERCOT), env=creds, clock=lambda: now)
    follow = keyed.follow_up(request, raw, now.date(), cov)
    api = [r for r in follow if r.meta.get("kind") == "api"]
    assert len(api) == 15 and all(ercot_api.HOST in r.url for r in api)
    assert keyed.skipped == []

    two_of_three = ercot.Connector(registry.get(ERCOT), env={}, clock=lambda: now)
    follow = two_of_three.follow_up(
        request, raw, now.date(), Coverage(downloads={"RTMLZHBSPP_2026": stamps[1:]})
    )
    assert [r.meta["report"] for r in follow] == ["RTMLZHBSPP_2025", "RTMLZHBSPP_2026"]


RT_2026 = "doclookupId=1287900508"


def _annual_with(**overrides) -> dict:
    """`ANNUAL` with some routes replaced (a URL substring -> fixture name or answer)."""
    return {**ANNUAL, **overrides}


def _workbook_calls(http: FakeHttp, doc: str = RT_2026) -> int:
    return sum(doc in c["url"] for c in http.calls)


def test_a_workbook_whose_parse_fails_is_downloaded_once_across_backfills(store, monkeypatch) -> None:
    """Review 2026-10-23, reproduced: a 2026 real-time workbook the parser rejects. Before the
    ledger, each of five daily backfills downloaded it again and recorded none of the downloads."""
    monkeypatch.setattr(ercot.Connector, "backfill_days", 5)
    now = dt.datetime(2026, 1, 3, 12, 0, tzinfo=dt.UTC)
    rejected = lambda **_k: (200, b"PK\x03\x04 not a workbook", {})  # noqa: E731
    total = 0
    for day in range(5):
        http = FakeHttp(_annual_with(**{RT_2026: rejected}))
        res = run(
            ERCOT, "backfill", store=store, http=http, now=now + dt.timedelta(days=day), trigger="backfill"
        )
        assert res.status == "failed" and "zip" in res.record["error"]
        total += _workbook_calls(http)
        # The run record names the download although its parse failed.
        entry = next(e for e in res.record["requests"] if e["key"].startswith("annual-RTMLZHBSPP_2026"))
        assert entry["downloaded"] is (day == 0) and entry["reused"] is (day > 0)
    assert total == 1
    ledger = [e for e in pstore.read_ledger(store, ERCOT) if e["report"] == "RTMLZHBSPP_2026"]
    assert len(ledger) == 1 and ledger[0]["state"] == "stored"
    assert len(pstore.backfill_attempts(store, ERCOT)) == 5


def test_a_run_killed_after_the_download_still_counts_it_and_reuses_the_bytes(store, monkeypatch) -> None:
    """An out-of-memory kill or a timeout ends the process with no run record; the ledger written
    before the parse still shows the download, and the next run reads the stored bytes."""
    monkeypatch.setattr(ercot.Connector, "backfill_days", 5)
    now = dt.datetime(2026, 1, 3, 12, 0, tzinfo=dt.UTC)
    real_frames = ercot.Connector.frames

    def killed(self, raw):
        if raw.meta.get("report") == "RTMLZHBSPP_2026":
            raise MemoryError("killed while parsing")
        yield from real_frames(self, raw)

    monkeypatch.setattr(ercot.Connector, "frames", killed)
    http = FakeHttp(ANNUAL)
    with pytest.raises(MemoryError):
        run(ERCOT, "backfill", store=store, http=http, now=now, trigger="backfill")
    assert store.runs(ERCOT) == [] and _workbook_calls(http) == 1
    assert pstore.backfill_attempts(store, ERCOT)  # the marker was written before the fetch

    monkeypatch.setattr(ercot.Connector, "frames", real_frames)
    http2 = FakeHttp(ANNUAL)
    res = run(ERCOT, "backfill", store=store, http=http2, now=now + dt.timedelta(hours=1), trigger="backfill")
    assert res.status == "ok" and _workbook_calls(http2) == 0
    entry = next(e for e in res.record["requests"] if e["key"].startswith("annual-RTMLZHBSPP_2026"))
    assert entry["reused"] and not entry["downloaded"]


def test_a_download_whose_outcome_was_never_seen_counts_against_the_cap(store, monkeypatch) -> None:
    """A process killed mid-download leaves its ledger entry at `started`: counted, so three such
    kills spend the year's three downloads and a fourth is refused without a request."""
    monkeypatch.setattr(ercot.Connector, "backfill_days", 5)
    now = dt.datetime(2026, 1, 3, 12, 0, tzinfo=dt.UTC)

    def dies(**_k):
        raise MemoryError("killed mid-download")

    for hour in range(3):
        with pytest.raises(MemoryError):
            run(
                ERCOT,
                "backfill",
                store=store,
                http=FakeHttp(_annual_with(**{RT_2026: dies})),
                now=now + dt.timedelta(hours=hour),
                trigger="backfill",
            )
    states = [e["state"] for e in pstore.read_ledger(store, ERCOT) if e["report"] == "RTMLZHBSPP_2026"]
    assert states == ["started"] * 3
    http = FakeHttp(ANNUAL)
    res = run(ERCOT, "backfill", store=store, http=http, now=now + dt.timedelta(hours=4), trigger="backfill")
    assert _workbook_calls(http) == 0
    assert any(s["report"] == "RTMLZHBSPP_2026" for s in res.record["skipped"])


def test_the_cap_holds_when_a_stored_workbook_fails_its_sha(store, monkeypatch) -> None:
    """The SHA-mismatch path fetches again only within the cap (review 2026-10-23)."""
    monkeypatch.setattr(ercot.Connector, "backfill_days", 5)
    now = dt.datetime(2026, 1, 3, 12, 0, tzinfo=dt.UTC)
    first = run(ERCOT, "backfill", store=store, http=FakeHttp(ANNUAL), now=now, trigger="backfill")
    entry = next(e for e in first.record["requests"] if e["key"].startswith("annual-RTMLZHBSPP_2026"))
    for i in range(2):  # two more downloads of the same report this year, from earlier runs
        pstore.write_ledger(
            store,
            ERCOT,
            f"2025121{i}T000000Z",
            entry["key"],
            {"state": "stored", "report": "RTMLZHBSPP_2026", "started_at": f"2025-12-1{i}T00:00:00+00:00"},
        )
    store.backend.put(entry["object_key"], b"corrupt")
    http = FakeHttp(ANNUAL)
    res = run(ERCOT, "backfill", store=store, http=http, now=now + dt.timedelta(hours=1), trigger="backfill")
    assert _workbook_calls(http) == 0
    refused = [s for s in res.record["skipped"] if s.get("report") == "RTMLZHBSPP_2026"]
    assert refused and "cap 3" in refused[0]["reason"]


def test_a_capped_download_is_sent_once(registry) -> None:
    """`PoliteSession` retries a 5xx up to four times inside one call; a capped report is asked once,
    and a 503 (no report served) is recorded as no download."""
    import requests

    from pipeline.connectors.http import HttpFailed, PoliteSession

    sent: list[str] = []

    def answer(method, url, **_kw):
        sent.append(url)
        resp = requests.models.Response()
        resp.status_code, resp._content, resp.url = 503, b"busy", url
        return resp

    http = PoliteSession(sleep=lambda _s: None)
    http.session.request = answer  # type: ignore[method-assign]
    c = ercot.Connector(registry.get(ERCOT), http=http, env={})
    url = "https://www.ercot.com/misdownload/servlets/mirDownload?doclookupId=1"
    capped = FetchRequest(key="annual-x", url=url, ext="zip", meta={"report": "X", "max_downloads": 3})
    with pytest.raises(HttpFailed):
        c.fetch(capped)
    assert len(sent) == 1 and http.max_attempts == 4  # restored for every other request
    with pytest.raises(HttpFailed):
        c.fetch(FetchRequest(key="list", url=url, ext="json"))
    assert len(sent) == 1 + 4
    assert not pstore.counts_as_download({"state": "failed", "http_status": 503})
    assert pstore.counts_as_download({"state": "failed", "http_status": None})  # a reset mid-body
    assert pstore.counts_as_download({"state": "started"})


def test_a_backfill_parses_only_its_window(store, monkeypatch) -> None:
    """A 5-day window from 2026-01-03 needs none of the 2025 workbooks' November sheet: those rows
    are not parsed or stored, and the workbook is still stored whole for a later, longer window."""
    monkeypatch.setattr(ercot.Connector, "backfill_days", 5)
    now = dt.datetime(2026, 1, 3, 12, 0, tzinfo=dt.UTC)
    res = run(ERCOT, "backfill", store=store, http=FakeHttp(ANNUAL), now=now, trigger="backfill")
    assert res.status == "ok" and res.record["rows_new"] == 2820 + 705
    assert "2025-11" not in pstore.list_months(store, ERCOT)
    rt_2025 = next(e for e in res.record["requests"] if e["key"].startswith("annual-RTMLZHBSPP_2025"))
    assert rt_2025["rows"] == 0 and rt_2025["bytes"] == len(fixture_bytes("ercot_rtm_annual_2025_nov02.zip"))
    assert ercot.sheet_filter(2025, dt.date(2025, 12, 29))("Nov") is False
    assert ercot.sheet_filter(2025, dt.date(2025, 12, 29))("Dec") is True
    assert ercot.sheet_filter(2025, dt.date(2025, 12, 29))("Sheet1") is True  # unknown: read it


def _isolate_everything(monkeypatch) -> None:
    from pipeline.connectors.prices import runner

    monkeypatch.setattr(ercot.Connector, "isolate_parse_bytes", 0)
    monkeypatch.setattr(runner, "ISOLATE_MERGE_ROWS", 1)


@pytest.mark.skipif(not can_fork(), reason="needs fork (Linux)")
def test_parsing_and_merging_in_child_processes_store_the_same_rows(tmp_path, monkeypatch) -> None:
    """Review 2026-10-23: the annual workbooks are parsed, and the backfill's months merged, in
    forked children whose memory goes back to the operating system; the store they leave is the
    in-process run's, row for row."""
    monkeypatch.setattr(ercot.Connector, "backfill_days", 70)
    now = dt.datetime(2026, 1, 3, 12, 0, tzinfo=dt.UTC)
    plain = Store(tmp_path / "plain")
    first = run(ERCOT, "backfill", store=plain, http=FakeHttp(ANNUAL), now=now, trigger="backfill")
    _isolate_everything(monkeypatch)
    forked = Store(tmp_path / "forked")
    second = run(ERCOT, "backfill", store=forked, http=FakeHttp(ANNUAL), now=now, trigger="backfill")
    assert first.status == second.status == "ok"
    for key in ("rows_new", "rows_seen", "rows_fetched", "stats", "dq_status"):
        assert first.record[key] == second.record[key], key
    assert pstore.list_months(plain, ERCOT) == pstore.list_months(forked, ERCOT)
    for month in pstore.list_months(plain, ERCOT):
        a = pstore.read_month(plain, ERCOT, month)
        b = pstore.read_month(forked, ERCOT, month)
        assert a is not None and b is not None
        pd.testing.assert_frame_equal(a, b, check_dtype=False)

    # Run again: the stored workbooks are re-read in children and every row is found unchanged.
    third = run(
        ERCOT,
        "backfill",
        store=forked,
        http=FakeHttp(ANNUAL),
        now=now + dt.timedelta(hours=1),
        trigger="backfill",
    )
    assert third.status == "unchanged" and third.record["rows_unchanged"] == first.record["rows_new"]


@pytest.mark.skipif(not can_fork(), reason="needs fork (Linux)")
def test_a_parse_child_killed_out_of_memory_fails_the_run_not_the_process(store, monkeypatch) -> None:
    """The out-of-memory killer takes the child; the run records the failure, its run record names
    the download, and the ledger has it stored for the next run to re-use."""
    import os
    import signal

    monkeypatch.setattr(ercot.Connector, "backfill_days", 5)
    _isolate_everything(monkeypatch)
    real_frames = ercot.Connector.frames

    def oom(self, raw):
        if raw.meta.get("report") == "RTMLZHBSPP_2026":
            os.kill(os.getpid(), signal.SIGKILL)
        yield from real_frames(self, raw)

    monkeypatch.setattr(ercot.Connector, "frames", oom)
    now = dt.datetime(2026, 1, 3, 12, 0, tzinfo=dt.UTC)
    res = run(ERCOT, "backfill", store=store, http=FakeHttp(ANNUAL), now=now, trigger="backfill")
    assert res.status == "failed" and res.record["error_class"] == "ChildFailed"
    assert "exit code -9" in res.record["error"]
    entry = next(
        e for e in store.runs(ERCOT)[-1]["requests"] if e["key"].startswith("annual-RTMLZHBSPP_2026")
    )
    assert entry["downloaded"] and entry["sha256"]
    ledger = [e for e in pstore.read_ledger(store, ERCOT) if e["report"] == "RTMLZHBSPP_2026"]
    assert [e["state"] for e in ledger] == ["stored"]
    assert pstore.list_months(store, ERCOT) == []  # nothing merged from a failed run


def test_a_child_s_exception_comes_back_as_one_the_runner_records() -> None:
    from pipeline.connectors.base import ParseError
    from pipeline.connectors.prices.runner import in_child

    def bad() -> None:
        raise ParseError("not a zip archive")

    with pytest.raises(ParseError, match="not a zip archive"):
        in_child(bad)
    assert in_child(lambda: {"rows": 3}) == {"rows": 3}


def test_the_api_is_used_only_with_all_three_settings() -> None:
    assert ercot.api_credentials({}) is None
    assert ercot.api_credentials({"ERCOT_API_SUBSCRIPTION_KEY": "k", "ERCOT_API_USERNAME": "u"}) is None
    assert (
        ercot.api_credentials(
            {"ERCOT_API_SUBSCRIPTION_KEY": "k", "ERCOT_API_USERNAME": "u", "ERCOT_API_PASSWORD": " "}
        )
        is None
    )
    assert ercot.api_credentials(dict.fromkeys(ercot.API_ENV, "v")) == dict.fromkeys(ercot.API_ENV, "v")


def test_daily_asks_for_tomorrow_s_day_ahead_page_and_tolerates_its_absence(store) -> None:
    http = FakeHttp(
        {
            "20261021_real_time_spp.html": "ercot_20261021_real_time_spp.html",
            "20261022_dam_spp.html": "ercot_20261022_dam_spp.html",
        }
    )
    res = run(ERCOT, "daily", store=store, http=http, now=NOW)
    asked = [c["url"].rsplit("/", 1)[-1] for c in http.calls]
    assert "20261023_dam_spp.html" in asked and "20261022_real_time_spp.html" not in asked
    tomorrow = next(e for e in res.record["requests"] if e["key"] == "da-20261023")
    assert tomorrow["http_status"] == 404 and "error" not in tomorrow
    stored = pstore.read_month(store, ERCOT, "2026-10")
    assert stored is not None
    assert set(stored["market"]) == {"rt", "da"}
    assert len(stored) == 96 * 15 + 24 * 15


def test_interval_months_split_on_utc_month_boundaries(store) -> None:
    http = FakeHttp({"system-prices/2026-10-21": "elexon_system_prices_20261021.json"})
    run(ELEXON, "daily", store=store, http=http, now=NOW)
    assert pstore.list_months(store, ELEXON) == ["2026-10"]
    df = pstore.read_month(store, ELEXON, "2026-10")
    assert df is not None
    assert df["interval_start"].min() == pd.Timestamp("2026-10-20T23:00:00Z")


# ----------------------------------------------------------------------------------------- plans
def _days(first: dt.date, last: dt.date) -> set[dt.date]:
    return {first + dt.timedelta(days=i) for i in range((last - first).days + 1)}


def test_nyiso_reads_days_past_the_daily_file_horizon_from_their_month_zip(registry) -> None:
    from pipeline.connectors.prices import nyiso

    c = nyiso.Connector(registry.get("us.iso.nyiso.prices"))
    today = dt.date(2026, 10, 22)
    keys = [r.key for r in c.plan("backfill", today, Coverage())]
    # 13 months per market (2025-10 .. 2026-10, the current month's zip included), no daily files.
    assert (
        sum(k.startswith("da-month-") for k in keys) == 13
        and sum(k.startswith("rt-month-") for k in keys) == 13
    )
    assert len(keys) == 26

    held = Coverage(
        complete_days={
            "da": _days(dt.date(2025, 10, 1), dt.date(2026, 10, 7)),
            "rt": _days(dt.date(2025, 10, 1), dt.date(2026, 10, 7)),
        }
    )
    # 10-08 .. 10-12 are past the horizon (their files are gone), so October's zip stands for them all.
    assert [r.key for r in c.plan("daily", today, held)] == ["da-month-202610", "rt-month-202610"]

    recent = Coverage(
        complete_days={
            "da": _days(dt.date(2026, 10, 1), dt.date(2026, 10, 20)),
            "rt": _days(dt.date(2026, 10, 1), dt.date(2026, 10, 20)),
        }
    )
    plan = c.plan("daily", today, recent)
    assert [r.key for r in plan] == ["da-20261021", "da-20261022", "da-20261023", "rt-20261021"]
    assert next(r for r in plan if r.key == "da-20261023").meta["optional"] is True


def test_caiso_asks_for_at_most_31_days_per_request(registry) -> None:
    from pipeline.connectors.prices import caiso

    c = caiso.Connector(registry.get("us.iso.caiso.oasis_prices"))
    plan = c.plan("backfill", dt.date(2026, 10, 22), Coverage())
    assert len(plan) == 24
    for r in plan:
        first, last = dt.date.fromisoformat(r.meta["first"]), dt.date.fromisoformat(r.meta["last"])
        assert (last - first).days < caiso.MAX_WINDOW_DAYS
        assert len(r.params["node"].split(",")) == 7
    assert caiso.oasis_time(dt.date(2026, 3, 8)) == "20260308T08:00-0000"
    assert caiso.oasis_time(dt.date(2026, 3, 9)) == "20260309T07:00-0000"


def test_spp_asks_little(registry) -> None:
    from pipeline.connectors.prices import spp

    c = spp.Connector(registry.get(SPP))
    today = dt.date(2026, 10, 22)
    assert len(c.plan("backfill", today, Coverage())) == 30
    assert [r.key for r in c.plan("latest", today, Coverage())] == ["rt-latest"]
    held = Coverage(complete_days={"da": _days(dt.date(2026, 10, 1), dt.date(2026, 10, 21))})
    assert [r.key for r in c.plan("daily", today, held)] == ["da-map"]


def test_elexon_latest_is_today_once(registry) -> None:
    from pipeline.connectors.prices import elexon

    c = elexon.Connector(registry.get(ELEXON))
    plan = c.plan("latest", dt.date(2026, 10, 22), Coverage())
    assert [r.url for r in plan] == [elexon.API.format(day="2026-10-22")]
    assert len(c.plan("backfill", dt.date(2026, 10, 22), Coverage())) == 366
