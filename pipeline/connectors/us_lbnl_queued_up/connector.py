"""us.lbnl.queued_up: Berkeley Lab "Queued Up" project-level interconnection queue data file.

What it is (read 2026-10-23, `docs/13` §2.1). Berkeley Lab, "in collaboration with GridTracker",
compiles the generation interconnection queues of the 7 ISOs/RTOs and 50 non-ISO balancing areas
into one workbook a year (2026 edition: requests through 2025, 38,201 rows). The publication page on
eta-publications.lbl.gov states: "The Queued Up data file is licensed CC BY 4.0. You may use, share,
or adapt the dataset as long as you attribute it to Lawrence Berkeley National Laboratory and
GridTracker." emp.lbl.gov, the project's home page, answers non-browser clients with a Cloudflare
block page; it is never fetched and nothing works around it. GridTracker is the company behind
Interconnection.fyi, a private aggregator this platform never scrapes (CLAUDE.md): this connector
reads only the file Berkeley Lab publishes under its own licence, from Berkeley Lab's host.

Held, not published (2026-10-23 review). The manifest row is `reuse: unknown`, `publication: none`,
the class of the rows rather than of LBNL's licence: the registry refuses scheduled runs (a person may
run the connector into the quarantine store with `--allow-restricted`), the loader refuses every row,
and nothing reaches the store, the resolver or a public page. Two reasons, both in `docs/13` §2.1:
(1) the kept rows restate the non-ISO utilities' own OASIS queues, whose terms
`data/sources.yaml` records as unknown (`us.oasis.non_iso_queues`), and the rule this module applies
to MISO (a compiler's CC BY cannot license what the operator's own terms withhold) applies to them
too until the owner decides otherwise; (2) 2,794 of the 2,914 pending requests carry no project name
and most no developer, so `pipeline/resolve.py` cannot score them against EIA-860M's planned units
(evidence 0.35 < `MIN_EVIDENCE`) and the same plant would show twice.

What is kept. Only the rows LBNL places in the non-ISO regions (`region` West or Southeast: TVA,
Southern, Duke, FPL, PacifiCorp, BPA, APS, NV Energy, Xcel Colorado and 40 more transmission
providers), the queues no other connector reads. Every other row is dropped in `parse`, never
stored as a record, and counted in the snapshot's `meta`:
- ERCOT, CAISO and NYISO are read first-hand (`us.iso.*.gen_queue`), fresher than an annual file,
  and so is SPP, whose own queue is published as derived fields under the owner's decision of
  2026-10-24 (`data/sources.yaml`; until its connector lands, SPP rows show from neither source);
- PJM rows are not public until a PJM licence exists (CLAUDE.md), whatever LBNL's licence says,
  and ISO-NE and MISO follow their operators' recorded status (`restricted`, `unknown` in
  `data/sources.yaml`), because a CC BY notice from a compiler cannot license what
  the operator's own terms withhold (the reading `docs/13` §6 applies to the LevelTen sheets in
  LBNL's solar and wind files).
A region label this module does not know is dropped and counted too, so a new label cannot publish
rows nobody classified.

Fetch. Only the edition a reviewer has read (`KNOWN_EDITION`) is ever fetched. The next editions'
pages are asked for first; if one answers 200 the run is `blocked` until a reviewer reads its
licence and records it, because CC BY requires the notice as supplied and the manifest's credit and
notice name the reviewed edition. The reviewed page's Notes block (the licence statement and
copyright line, with its links) must hash to `NOTES_SHA256`, and the two licence sentences must be
on it, else the run is `blocked` (an added restriction, not only a removed sentence, stops the run).
Then the workbook it links, conditional on the previous snapshot's validators. robots.txt on
eta-publications.lbl.gov allows both paths (read 2026-10-23).

Cleaning in `parse` and `normalize`:
- Repeated positions. `<entity>/<q_id>` is the record key ("Combine with entity to form a unique
  identifier across the full dataset", codebook). The 2026 edition repeats 31 positions in the kept
  regions, TVA 29 and Duke Energy Florida 2 (Hines, Polk County): each is one request listed twice,
  once with its county and state and once without. A copy that agrees with another row of the same
  position on `q_date`, `utility`, `type_clean` and `mw_1`, and on every other field both carry, is
  dropped in favour of the row that carries more fields (the located one), and counted in
  `meta.repeats_collapsed`. A repeat that disagrees (Duke Energy Florida's position 2: 1,701 MW
  located, 701 MW not) stays two records with content-derived `#h...` suffixes
  (`dedupe_strategy = "suffix"`) and a DQ warning; the one "not assigned" position is a content hash.
- Placeholders. LBNL's `developer` and `project_name` carry "N/A" (863 rows), "Masked" (203) and
  "Third Party" (28) where no name is given; these, and the other spellings in `PLACEHOLDER_NAMES`,
  are no name (None), never a company or a project.
- Natural persons. A `developer` that `services.personal_names` reads as a person's name (6 names,
  7 rows in the 2026 edition) is removed from the row in `parse`, so neither the sponsor nor the
  `raw` payload holds it (CLAUDE.md: store the minimum personal data; the snapshot file is LBNL's
  own, unedited). The count is in `meta.developers_withheld_as_persons`.
- Operator. `iso` is the transmission provider's name, not LBNL's `entity` code ("SRP_ANPP",
  "WAPA-MPP", one "Duke" for three Duke utilities): `utility`, except where that column names the
  interconnection customer (`ENTITY_OPERATORS`: BPA), else the one utility the entity's other rows
  name, else `ENTITY_FALLBACK`, else the code.
Lifecycle: `status_map.yaml` beside this file. Capacity: `mw_1`, the first resource type's MW (a
hybrid's storage MW stays in `raw` as `mw_2`); LBNL's imputed storage figures are not in the file.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import html
import io
import json
import pathlib
import re
import time
from collections import Counter, defaultdict
from typing import Any, ClassVar, Literal
from urllib.parse import urljoin

import openpyxl
import pandas as pd

from pipeline.connectors.base import BlockedError, ConnectorError, Kind, ParseError, RawSnapshot, content_hash
from pipeline.connectors.base import Connector as BaseConnector
from pipeline.connectors.canonical import (
    classify_tech,
    harmonise_status,
    norm_county,
    norm_name,
    norm_org,
    norm_state,
    to_date,
    to_float,
)
from pipeline.connectors.dedupe import split_key
from pipeline.connectors.http import HttpBlocked
from services.personal_names import is_personal_name

EDITION_PAGE = "https://eta-publications.lbl.gov/publications/queued-{year}-edition-characteristics"
#: The edition a reviewer read and recorded (2026-10-23): requests through 2025, published 05/2026.
#: Moving it is a review step: read the new page, record its licence and notice in
#: data/sources.yaml (`license`, `attribution`, `notice`) and docs/13 §2.1, and set NOTES_SHA256.
KNOWN_EDITION = 2026
#: sha256 of `notes_digest_text` over KNOWN_EDITION's publication page, read 2026-10-23
#: (tests/fixtures/lbnl_queued_up_publication.html holds that row of the page, markup unedited).
NOTES_SHA256 = "9cda6bb021be0b2f0beee74c3b918f46aa0cb138f2254bcc161accb386255c19"
XLSX_MAGIC = b"PK"
XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

#: The licence sentences the publication page carried on 2026-10-23, whitespace-normalised. Each
#: must still be on the page, and the page must link the CC BY 4.0 deed, or nothing is fetched.
LICENCE_SENTENCES = (
    "The Queued Up data file is licensed CC BY 4.0.",
    "You may use, share, or adapt the dataset as long as you attribute it to Lawrence Berkeley "
    "National Laboratory and GridTracker.",
)
LICENCE_DEED = "creativecommons.org/licenses/by/4.0"

#: LBNL's `region` values whose rows this connector keeps: the non-ISO transmission providers.
PUBLISHED_REGIONS = frozenset({"West", "Southeast"})
#: Regions read first-hand by another connector.
FIRST_HAND_REGIONS = {
    "ERCOT": "us.iso.ercot.gen_queue",
    "CAISO": "us.iso.caiso.gen_queue",
    "NYISO": "us.iso.nyiso.gen_queue",
    "SPP": "us.iso.spp.gen_queue",
}
#: Regions whose operators' terms withhold their rows (module docstring).
WITHHELD_REGIONS = {
    "PJM": "us.iso.pjm.gen_queue",
    "MISO": "us.iso.miso.gen_queue",
    "ISO-NE": "us.iso.isone.gen_queue",
}
#: The register row whose unread terms hold this source's own kept rows (module docstring).
OPERATOR_TERMS_SOURCE = "us.oasis.non_iso_queues"
QUEUE_SHEET_SUFFIX = "complete queue data"
#: Columns the normaliser reads; a workbook without any of them fails the parse closed.
REQUIRED_COLUMNS = (
    "q_id",
    "q_status",
    "q_date",
    "prop_date",
    "on_date",
    "wd_date",
    "ia_date",
    "IA_phase_clean",
    "county",
    "state",
    "region",
    "project_name",
    "utility",
    "entity",
    "developer",
    "type_clean",
    "mw_1",
)
#: Fields two listings of one position must share (blank on both counts as shared) to be one row.
REPEAT_IDENTITY = ("q_date", "utility", "type_clean", "mw_1")
#: Spellings LBNL (or the utility it read) uses where no name is given, compared case-folded with
#: hyphens read as spaces and outer dots dropped. Never a company or project name.
PLACEHOLDER_NAMES = frozenset(
    {
        "n/a",
        "na",
        "none",
        "null",
        "masked",
        "third party",
        "tbd",
        "tba",
        "unknown",
        "confidential",
        "withheld",
        "redacted",
        "not available",
        "not applicable",
        "not provided",
    }
)
#: LBNL `entity` codes whose `utility` column names the interconnection customer, not the
#: transmission provider (BPA, 2026 edition: 780 rows blank, the other 218 naming 115 different customers).
ENTITY_OPERATORS = {"BPA": "Bonneville Power Administration"}
#: The operator name for a row with no `utility` whose entity's other rows name several utilities.
ENTITY_FALLBACK = {"Duke": "Duke Energy"}
_TAG = re.compile(r"<[^>]+>")
_WS = re.compile(r"\s+")
_HREF = re.compile(r'\bhref="([^"]*)"', re.I)
_XLSX_LINK = re.compile(r'<a\s[^>]*href="([^"]+\.xlsx)"[^>]*>(.*?)</a>', re.I | re.S)
_NOTES_ROW = re.compile(r"<tr\b[^>]*\bbibcite-notes\b[^>]*>(.*?)</tr>", re.I | re.S)


def page_text(markup: str) -> str:
    """The page's text with markup removed, entities decoded and whitespace collapsed."""
    return _WS.sub(" ", html.unescape(_TAG.sub(" ", markup)).replace("\xa0", " ")).strip()


def licence_problems(markup: str) -> list[str]:
    """What is missing from the publication page's licence statement (empty: the licence stands)."""
    text = page_text(markup).replace(" .", ".").replace(" ,", ",")
    missing = [s for s in LICENCE_SENTENCES if s not in text]
    if LICENCE_DEED not in markup:
        missing.append(f"a link to {LICENCE_DEED}")
    return missing


def notes_digest_text(markup: str) -> str | None:
    """The publication page's Notes row as compared: its text (`page_text`), then each link it
    holds, one per line. None when the page has no Notes row."""
    blocks = _NOTES_ROW.findall(markup)
    if len(blocks) != 1:
        return None
    block = blocks[0]
    return "\n".join([page_text(block), *(html.unescape(h) for h in _HREF.findall(block))])


def notes_sha256(markup: str) -> str | None:
    text = notes_digest_text(markup)
    return None if text is None else hashlib.sha256(text.encode("utf-8")).hexdigest()


def data_file_url(markup: str, page_url: str) -> str | None:
    """The "Data File XLSX" link on the publication page (else its first xlsx link), absolute."""
    links = [(href, page_text(label)) for href, label in _XLSX_LINK.findall(markup)]
    if not links:
        return None
    chosen = next((href for href, label in links if "data file" in label.lower()), links[0][0])
    return urljoin(page_url, html.unescape(chosen))


def newer_editions(now: dt.datetime, limit: int = 3) -> list[int]:
    """The editions after the reviewed one that may have been published by `now` (an edition names
    the year it is published in), oldest first, at most `limit`. Any of them answering is a new
    edition nobody has reviewed."""
    return list(range(KNOWN_EDITION + 1, min(now.year + 1, KNOWN_EDITION + limit) + 1))


def _blank(v: Any) -> bool:
    if v is None:
        return True
    if isinstance(v, float) and pd.isna(v):
        return True
    return isinstance(v, str) and not v.strip()


def _text(v: Any) -> str | None:
    if _blank(v):
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return _WS.sub(" ", str(v)).strip()


def name_or_none(v: Any) -> str | None:
    """A developer or project name, or None for a blank or a placeholder (`PLACEHOLDER_NAMES`)."""
    text = _text(v)
    if text is None:
        return None
    key = _WS.sub(" ", text.casefold().replace("-", " ")).strip(" .")
    return None if not key or key in PLACEHOLDER_NAMES else text


def record_key(row: dict[str, Any]) -> str:
    """`<entity>/<q_id>`; a content hash when LBNL gives no position ("not assigned")."""
    entity, qid = _text(row.get("entity")) or "", _text(row.get("q_id"))
    if qid is None or qid.lower() == "not assigned":
        return f"{entity}/" + content_hash(
            entity,
            row.get("q_date"),
            row.get("county"),
            row.get("state"),
            row.get("mw_1"),
            row.get("type_clean"),
        )
    return f"{entity}/{qid}"


def _same(a: Any, b: Any) -> bool:
    if _blank(a) and _blank(b):
        return True
    if _blank(a) or _blank(b):
        return False
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return float(a) == float(b)
    return _text(a) == _text(b)


def is_copy_of(row: dict[str, Any], other: dict[str, Any]) -> bool:
    """Whether two listings of one position are one request: the same `REPEAT_IDENTITY` fields and
    no field that both carry with different values."""
    if not all(_same(row.get(f), other.get(f)) for f in REPEAT_IDENTITY):
        return False
    return all(
        _same(row.get(f), other.get(f))
        for f in set(row) | set(other)
        if not _blank(row.get(f)) and not _blank(other.get(f))
    )


def collapse_repeats(rows: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], int]:
    """`rows` without the copies of a repeated position (module docstring), in their file order,
    and how many copies were dropped. Of a set of copies the row carrying the most fields stays."""
    by_key: dict[str, list[int]] = defaultdict(list)
    for i, row in enumerate(rows):
        by_key[record_key(row)].append(i)
    dropped: set[int] = set()
    for members in by_key.values():
        if len(members) < 2:
            continue
        kept: list[int] = []
        for i in sorted(members, key=lambda i: (-sum(not _blank(v) for v in rows[i].values()), i)):
            if any(is_copy_of(rows[i], rows[k]) for k in kept):
                dropped.add(i)
            else:
                kept.append(i)
    return [row for i, row in enumerate(rows) if i not in dropped], len(dropped)


def operator_names(rows: list[dict[str, Any]]) -> list[str | None]:
    """The transmission provider each row's request is made to (module docstring, "Operator")."""
    utilities: dict[str, set[str]] = defaultdict(set)
    for row in rows:
        entity, utility = _text(row.get("entity")), _text(row.get("utility"))
        if entity and utility:
            utilities[entity].add(utility)
    out: list[str | None] = []
    for row in rows:
        entity, utility = _text(row.get("entity")), _text(row.get("utility"))
        if entity in ENTITY_OPERATORS:
            out.append(ENTITY_OPERATORS[entity])
        elif utility:
            out.append(utility)
        elif entity and len(utilities.get(entity, ())) == 1:
            out.append(next(iter(utilities[entity])))
        else:
            out.append(ENTITY_FALLBACK.get(entity or "", entity))
    return out


def technology_of(type_clean: Any) -> tuple[str, str]:
    """(technology, kind) from LBNL's `type_clean` ("Solar+Battery", "Gas+Oil"): its `+`-joined
    components go through the shared component rule, so firm generation with storage is the
    generation and solar or wind with storage keeps its hybrid class."""
    text = _text(type_clean)
    if text is None:
        return classify_tech(None)
    return classify_tech(" + ".join(p.strip() for p in text.split("+") if p.strip()))


class Connector(BaseConnector):
    source_id: ClassVar[str] = "us.lbnl.queued_up"
    kind: ClassVar[Kind] = "proposal"
    ext: ClassVar[str] = "xlsx"
    status_key: ClassVar[str] = "lbnl_queued_up"
    status_map_path: ClassVar[pathlib.Path | None] = pathlib.Path(__file__).with_name("status_map.yaml")
    dedupe_strategy: ClassVar[Literal["hold", "suffix"]] = "suffix"
    dq_required_fields: ClassVar[tuple[str, ...]] = ("capacity_mw", "technology_raw")
    key_source_columns: ClassVar[tuple[str, ...]] = (
        "q_id",
        "q_status",
        "entity",
        "utility",
        "region",
        "state",
        "county",
        "type_clean",
        "mw_1",
    )
    #: How many editions after KNOWN_EDITION one run asks for (`newer_editions`).
    max_editions: ClassVar[int] = 3

    @classmethod
    def project_root(cls, record_key: str) -> str:
        """The queue position without the `#h...` suffix a repeated position carries."""
        return split_key(record_key)[0]

    # ------------------------------------------------------------------ fetch
    def _get_page(self, url: str) -> Any:
        try:
            return self.http.get(url, timeout=60)
        except HttpBlocked as exc:
            raise BlockedError(f"{url}: {exc}") from exc

    def _publication_page(self) -> tuple[str, str, list[dict[str, Any]]]:
        tried: list[dict[str, Any]] = []
        for year in newer_editions(self.now(), self.max_editions):
            url = EDITION_PAGE.format(year=year)
            r = self._get_page(url)
            tried.append({"url": url, "status": r.status_code})
            if r.status_code == 404:
                continue
            if r.status_code == 200:
                raise BlockedError(
                    f"{url} answers 200: Queued Up's {year} edition is published, and only the "
                    f"{KNOWN_EDITION} edition has been reviewed. Nothing is fetched until a reviewer reads "
                    "the new page, records its licence, credit and notice in data/sources.yaml and docs/13 "
                    "§2.1, and moves KNOWN_EDITION and NOTES_SHA256 in pipeline/connectors/us_lbnl_queued_up"
                )
            raise ConnectorError(f"GET {url} -> HTTP {r.status_code}")
        url = EDITION_PAGE.format(year=KNOWN_EDITION)
        r = self._get_page(url)
        tried.append({"url": url, "status": r.status_code})
        if r.status_code != 200:
            raise ConnectorError(f"GET {url} -> HTTP {r.status_code}; pages asked: {tried}")
        return url, r.text, tried

    def _conditional_headers(self, data_url: str) -> dict[str, str]:
        if self.previous is None or self.previous.meta.get("data_url") != data_url:
            return {}
        meta = self.previous.meta
        headers: dict[str, str] = {}
        if meta.get("etag"):
            headers["If-None-Match"] = str(meta["etag"])
        if meta.get("last_modified"):
            headers["If-Modified-Since"] = str(meta["last_modified"])
        return headers

    def _stored_workbook(self) -> bytes | None:
        if self.previous is None:
            return None
        try:
            body = self.previous.content()
        except Exception:  # an unusable stored object costs one download, no more
            return None
        return body if body.startswith(XLSX_MAGIC) else None

    def fetch(self) -> RawSnapshot:
        t0 = time.monotonic()
        page_url, markup, tried = self._publication_page()
        missing = licence_problems(markup)
        if missing:
            raise BlockedError(
                f"{page_url}: the licence statement changed (missing {missing}); nothing is fetched until "
                "a reviewer re-reads it and updates data/sources.yaml and docs/13 §2.1"
            )
        digest = notes_sha256(markup)
        if digest != NOTES_SHA256:
            raise BlockedError(
                f"{page_url}: the Notes block (licence statement and copyright line) is not the one reviewed "
                f"on 2026-10-23 (sha256 {digest}, reviewed {NOTES_SHA256}); nothing is fetched until a "
                "reviewer re-reads it, updates data/sources.yaml and docs/13 §2.1 and sets NOTES_SHA256"
            )
        data_url = data_file_url(markup, page_url)
        if data_url is None:
            raise ConnectorError(f"{page_url} links no xlsx data file")
        conditional = self._conditional_headers(data_url)
        requests_made = len(tried) + 1
        r = (
            self.http.get(data_url, timeout=300, headers=conditional)
            if conditional
            else self.http.get(data_url, timeout=300)
        )
        content: bytes | None = None
        not_modified = False
        if r.status_code == 304 and conditional:
            content = self._stored_workbook()
            not_modified = content is not None
            if content is None:
                r = self.http.get(data_url, timeout=300)
                requests_made += 1
        if content is None:
            if r.status_code != 200:
                raise ConnectorError(f"GET {data_url} -> HTTP {r.status_code}")
            content = r.content
        previous = self.previous.meta if (self.previous is not None and not_modified) else {}
        etag = r.headers.get("ETag") or previous.get("etag")
        last_modified = r.headers.get("Last-Modified") or previous.get("last_modified")
        meta: dict[str, Any] = {
            "publication_page": page_url,
            "edition": KNOWN_EDITION,
            "pages_tried": tried,
            "data_url": data_url,
            "licence_checked": list(LICENCE_SENTENCES),
            "notes_sha256": digest,
        }
        if etag:
            meta["etag"] = str(etag)
        if last_modified:
            meta["last_modified"] = str(last_modified)
        if not_modified:
            meta["not_modified"] = True
        return RawSnapshot(
            content=content,
            content_type=XLSX_CONTENT_TYPE if not_modified else r.headers.get("Content-Type", ""),
            url=data_url,
            retrieved_at=dt.datetime.now(dt.UTC),
            http_status=r.status_code,
            ext="xlsx",
            headers=dict(r.headers),
            elapsed_s=round(time.monotonic() - t0, 2),
            requests_made=requests_made,
            meta=meta,
        )

    # ------------------------------------------------------------------ parse
    def parse(self, raw: RawSnapshot) -> list[dict[str, Any]]:
        if not raw.content.startswith(XLSX_MAGIC):
            raise ParseError(f"{raw.url} is not an xlsx (first bytes {raw.content[:16]!r})")
        try:
            wb = openpyxl.load_workbook(io.BytesIO(raw.content), read_only=True, data_only=True)
        except Exception as exc:  # a corrupt or truncated workbook
            raise ParseError(f"{raw.url}: unreadable workbook ({exc!r})") from exc
        try:
            sheet = next((ws for ws in wb.worksheets if ws.title.lower().endswith(QUEUE_SHEET_SUFFIX)), None)
            if sheet is None:
                raise ParseError(f"{raw.url}: no '{QUEUE_SHEET_SUFFIX}' sheet among {wb.sheetnames}")
            rows_iter = sheet.iter_rows(values_only=True)
            header: list[str] | None = None
            for i, row in enumerate(rows_iter):
                if row and _text(row[0]) == "q_id":
                    header = [str(c).strip() if c is not None else "" for c in row]
                    break
                if i >= 10:
                    break
            if header is None:
                raise ParseError(f"{raw.url}: sheet {sheet.title!r} has no q_id header row in its first rows")
            missing = [c for c in REQUIRED_COLUMNS if c not in header]
            if missing:
                raise ParseError(f"{raw.url}: sheet {sheet.title!r} lacks columns {missing}")
            kept: list[dict[str, Any]] = []
            regions: Counter[str] = Counter()
            for values in rows_iter:
                if values is None or all(_blank(v) for v in values):
                    continue
                row = {h: v for h, v in zip(header, values, strict=False) if h}
                region = _text(row.get("region")) or ""
                regions[region] += 1
                if region in PUBLISHED_REGIONS:
                    kept.append(row)
        finally:
            wb.close()
        if not kept:
            raise ParseError(
                f"{raw.url}: no row in the regions this connector publishes {sorted(PUBLISHED_REGIONS)}"
            )
        kept, collapsed = collapse_repeats(kept)
        withheld = 0
        for row in kept:
            developer = name_or_none(row.get("developer"))
            if developer is not None and is_personal_name(developer):
                row["developer"] = None
                withheld += 1
        raw.meta["rows_by_region"] = dict(sorted(regions.items()))
        raw.meta["rows_kept"] = len(kept)
        raw.meta["repeats_collapsed"] = collapsed
        raw.meta["developers_withheld_as_persons"] = withheld
        raw.meta["dropped_first_hand"] = {r: regions[r] for r in FIRST_HAND_REGIONS if regions[r]}
        raw.meta["dropped_withheld"] = {r: regions[r] for r in WITHHELD_REGIONS if regions[r]}
        known = PUBLISHED_REGIONS | set(FIRST_HAND_REGIONS) | set(WITHHELD_REGIONS)
        raw.meta["dropped_unrecognised_region"] = {r: n for r, n in regions.items() if r not in known}
        return kept

    # ------------------------------------------------------------------ normalise
    def _lifecycle(self, row: dict[str, Any]) -> tuple[str, str]:
        ctx = {"status_raw": _text(row.get("q_status")), "ia_phase": _text(row.get("IA_phase_clean"))}
        return harmonise_status(self.status_key, ctx, self.status_map)

    def normalize(self, rows: list[dict[str, Any]], raw: RawSnapshot) -> pd.DataFrame:
        tech = [technology_of(r.get("type_clean")) for r in rows]
        lifecycle = [self._lifecycle(r) for r in rows]
        names = [name_or_none(r.get("project_name")) for r in rows]
        sponsors = [name_or_none(r.get("developer")) for r in rows]
        mw = [to_float(r.get("mw_1")) for r in rows]
        df = pd.DataFrame(
            {
                "source_record_id": [record_key(r) for r in rows],
                "kind": [k for _, k in tech],
                "name_canonical": names,
                "name_norm": [norm_name(n) for n in names],
                "sponsor_name": sponsors,
                "sponsor_norm": [norm_org(s) for s in sponsors],
                "technology": [t for t, _ in tech],
                "technology_raw": [_text(r.get("type_clean")) for r in rows],
                "capacity_mw": pd.array(
                    [m if m is not None and m > 0 else None for m in mw], dtype="Float64"
                ),
                "storage_mwh": pd.array([None] * len(rows), dtype="Float64"),
                # The transmission provider by name (Tennessee Valley Authority, Duke Energy
                # Florida): what the proposal page prints under "ISO / operator" for a non-ISO queue.
                "iso": operator_names(rows),
                "state": [norm_state(r.get("state")) for r in rows],
                "county": [_text(r.get("county")) for r in rows],
                "county_norm": [norm_county(r.get("county")) for r in rows],
                "lifecycle_state": [s for s, _ in lifecycle],
                "status_raw": [_text(r.get("q_status")) for r in rows],
                "status_rule": [rule for _, rule in lifecycle],
                "status_conflict": False,
                "queue_date": [to_date(r.get("q_date")) for r in rows],
                "proposed_cod": [to_date(r.get("prop_date")) for r in rows],
                "study_phase": [_text(r.get("IA_phase_clean")) for r in rows],
                "ia_date": [to_date(r.get("ia_date")) for r in rows],
                "withdrawn_date": [to_date(r.get("wd_date")) for r in rows],
                "actual_cod": [to_date(r.get("on_date")) for r in rows],
                "queue_id": [_text(r.get("q_id")) for r in rows],
                "eia_plant_id": pd.array([None] * len(rows), dtype="string"),
                "eia_generator_id": pd.array([None] * len(rows), dtype="string"),
                "cross_refs": "",
            },
            index=range(len(rows)),
        )
        for col in ("queue_date", "proposed_cod", "ia_date", "withdrawn_date", "actual_cod"):
            df[col] = pd.to_datetime(df[col], errors="coerce")
        return self.finalize(df, rows, raw)

    def restate_status(self, df: pd.DataFrame) -> pd.DataFrame | None:
        """Re-harmonise each stored row's own `raw` payload under the current status map."""
        if "raw" not in df.columns:
            return None
        rows = [json.loads(r) if isinstance(r, str) and r else {} for r in df["raw"]]
        lifecycle = [self._lifecycle(r) for r in rows]
        return pd.DataFrame(
            {"lifecycle_state": [s for s, _ in lifecycle], "status_rule": [r for _, r in lifecycle]},
            index=df.index,
        )
