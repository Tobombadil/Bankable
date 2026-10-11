"""The owner's 100-record hand check of sites and merged records (docs/21 §3.25 "Hand check";
decisions log 2026-10-10).

    python -m services.sites.handcheck generate --database-url URL --out DIR [--seed N] [--posture P]
    python -m services.sites.handcheck score FILLED.xlsx [--json]

The check decides two things: whether sites stay on for the beta (more than about one site in ten
grouping records that are not the same project switches them off, `SITES_ENABLED=0`), and whether
merged records need splitting (the claim "one record per project").

**`generate`** reads a store **read-only** (`SET TRANSACTION READ ONLY` on Postgres, `PRAGMA
query_only` on SQLite; the session is rolled back) and writes a worksheet (`handcheck.xlsx`), a CSV
twin of each data sheet, `manifest.json` and a README. Every value comes from the served view at one
tier (`public` by default), through the API's own functions: `services.api.sites.served_site` for a
site and its labels, `services.api.visibility.gated_record` for each record's fields and sponsor,
`visible_source_links` and `services.api.proposal_members.member_row` for each source link (a
source record id the licence does not let the API print is replaced by a hash of it), and the
interconnection-point visibility predicate for connection-point names. Items that name a natural
person (a sponsor flagged by the personal-data pass, or any printed name the same classifier
reads as a person's, `services.personal_names`) are left out of the frame and counted.

The sample is 50 sites and 50 merged records by default, stratified and drawn with permanent random
numbers (`handcheck_stats.prn`), so the same seed draws the same items from the same store and, as
far as the stores agree, from another store too. Public ids are minted at load, so every row also
carries stable keys: `source_id:source_record_id` for each source link (`source_id:#<hash>` where
the id is not served) and the EIA plant ids. Links are built from the base-URL cell on the
Instructions sheet: empty gives relative paths.

* **Sites**, one stratum per item, in this order: the `--largest` sites by member count (checked in
  full); sites holding a `same_site` label (grouped, relation unclear: low confidence); then by the
  site's weakest grouping rule (`eia_plant`, `exact_point_*`, `poi_sponsor`, `poi_stem`) crossed
  with its lowest label confidence (`high`, `medium`, `low`).
* **Merged records** (a live record served with two or more source links), the first category
  that holds: several EIA plants in one record; several requests from one register; link names
  with differing phase markers; several generators of one EIA plant; one project in two or more
  registers.

**`score`** reads the filled workbook (or the directory of CSVs) and prints the site error rate with
a Wilson 95 % interval, the weighted estimate for all served sites, whether the 1-in-10 rule is
crossed on the point estimate and on the interval's lower bound, and the merged-record error rate
by stratum and weighted (`handcheck_stats`).
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
import os
import pathlib
import sys
from collections import Counter
from collections.abc import Callable, Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from typing import Any, Generic, TypeVar

from sqlalchemy import func, select, text
from sqlalchemy.orm import Session, selectinload

from services.db.models import Organization, Proposal, ProposalSource, Site
from services.labels import lifecycle_label, technology_label
from services.personal_names import is_personal_name, normalise
from services.sites import evidence, rules
from services.sites.handcheck_stats import (
    KILL_SWITCH_RATE,
    Estimate,
    Judged,
    allocate,
    decisive_counts,
    estimate,
    prn,
    wilson,
)

T = TypeVar("T")
ItemT = TypeVar("ItemT")

SITE_SAMPLE = 50
MERGED_SAMPLE = 50
#: The largest sites by member count are checked in full (a certainty stratum).
LARGEST_SITES = 5
MIN_PER_STRATUM = 2
DEFAULT_SEED = 20261014
WORKBOOK = "handcheck.xlsx"
#: Members shown per plant group before the rest are summarised in one line (all are on "Site members").
UNITS_SHOWN = 3

# Sheet names and the headers the scorer reads back.
SHEET_INSTRUCTIONS = "Instructions"
SHEET_SITES = "Sites"
SHEET_MERGED = "Merged records"
SHEET_SITE_MEMBERS = "Site members"
SHEET_MERGED_LINKS = "Merged links"
SHEET_STRATA = "Strata"
H_ITEM = "Item"
H_SITE_VERDICT = "All these belong to one project or site?"
H_SITE_PROBLEM = "If no, what is wrong"
H_MERGED_VERDICT = "One real project?"
H_MERGED_KIND = "What kind"
H_NOTES = "Notes"
H_STRATUM = "Stratum"
H_POPULATION = "Stratum population"
H_SAMPLED = "Stratum sample"

SITE_VERDICTS = ("Yes", "No", "Unsure")
SITE_PROBLEMS = ("Unrelated record grouped", "Missing record", "Wrong lead", "Wrong label", "Other")
#: A "No" with one of these problems (or none given) is a grouping error: it counts toward the
#: 1-in-10 rule. A missing record, a wrong lead or a wrong label is a quality issue of a site whose
#: members do belong together, reported apart.
GROUPING_PROBLEMS = frozenset({"unrelated record grouped", "other", ""})
MERGED_VERDICTS = ("Yes", "No (several projects)", "Unsure")
MERGED_KINDS = (
    "Phases of one development",
    "Separate units or plants at one site",
    "Re-filing or replacement",
    "Unrelated projects",
    "Other",
)

RULE_GROUP = {
    "eia_plant": "eia_plant",
    "exact_point_sponsor": "exact_point",
    "exact_point_stem": "exact_point",
    "poi_sponsor": "poi_sponsor",
    "poi_stem": "poi_stem",
}
RULE_TEXT = {
    "eia_plant": "Same EIA plant",
    "exact_point_sponsor": "Same exact location and same sponsor",
    "exact_point_stem": "Same exact location and same project name",
    "poi_sponsor": "Same grid connection point and same sponsor",
    "poi_stem": "Same grid connection point and same project name",
}
CONFIDENCE_ORDER = ("low", "medium", "high")
RELATION_TEXT = {
    "lead": "LEAD",
    "unit_of": "unit of #{parent}",
    "phase_of": "phase of the lead",
    "co_located": "co-located with the lead",
    "refiling_of": "re-filing of the lead",
    "superseded_by": "superseded by the lead",
    "expansion_of": "expansion of the lead",
    "expanded_by": "expanded by the lead",
    "same_site": "same site, relation unclear",
}
SITE_STRATA_TEXT = {
    "largest": "The largest sites by member count, all checked",
    "same_site": "Sites holding a 'same site, relation unclear' label (low confidence)",
}
GROUP_TEXT = {
    "eia_plant": "same EIA plant",
    "exact_point": "same exact location plus sponsor or name",
    "poi_sponsor": "same connection point and sponsor",
    "poi_stem": "same connection point and project name",
}
#: Merged-record categories, in the order the first that holds is taken.
MERGED_CATEGORIES: tuple[tuple[str, str], ...] = (
    ("multi_plant", "Several EIA plants in one record"),
    ("multi_request", "Several queue requests from one register"),
    ("phase_markers", "Source names with differing phase markers"),
    ("multi_generator", "Several generators of one EIA plant"),
    ("cross_register", "One project seen in two or more registers"),
)
MERGED_TEXT = dict(MERGED_CATEGORIES)


# ------------------------------------------------------------------------------------- items
@dataclass(frozen=True)
class Link:
    """One source link as the API serves it (`member_row`)."""

    source_id: str
    source_name: str
    key: str
    record_id: str | None
    name: str | None
    capacity_mw: float | None
    technology: str | None
    lifecycle_state: str | None
    role: str
    source_url: str

    @property
    def plant_id(self) -> str | None:
        return evidence.eia_plant_id(self.source_id, self.record_id) if self.record_id else None


@dataclass(frozen=True)
class Member:
    """One served site member, in served display order (`rank` 0 is the lead)."""

    rank: int
    public_id: str
    slug: str
    name: str
    capacity_mw: float | None
    technology: str | None
    lifecycle_state: str
    jurisdiction: str | None
    sponsor: str | None
    relation: str
    relation_rule: str
    confidence: str
    grouping_rule: str
    group_key: str
    parent_rank: int | None
    connection_point: str | None
    links: tuple[Link, ...]

    @property
    def plant_ids(self) -> tuple[str, ...]:
        return _plants(self.links)


@dataclass
class SiteItem:
    public_id: str
    name: str
    members: list[Member]
    stratum: str = ""

    @property
    def key(self) -> str:
        return min((link.key for m in self.members for link in m.links), default=self.public_id)

    @property
    def weakest_rule(self) -> str:
        return max((m.grouping_rule for m in self.members), key=_rule_index)

    @property
    def lowest_confidence(self) -> str:
        return min((m.confidence for m in self.members), key=_confidence_index)

    @property
    def plant_ids(self) -> tuple[str, ...]:
        return _plants(link for m in self.members for link in m.links)


@dataclass
class MergedItem:
    public_id: str
    slug: str
    name: str
    capacity_mw: float | None
    technology: str | None
    lifecycle_state: str
    jurisdiction: str | None
    sponsor: str | None
    links: tuple[Link, ...]
    flags: tuple[str, ...]
    site_public_id: str | None = None
    stratum: str = ""

    @property
    def key(self) -> str:
        return min((link.key for link in self.links), default=self.public_id)

    @property
    def plant_ids(self) -> tuple[str, ...]:
        return _plants(self.links)


@dataclass
class Frame(Generic[ItemT]):
    """The served items a sample is drawn from, and what was left out and why."""

    items: list[ItemT]
    excluded: Counter[str] = field(default_factory=Counter)
    #: Counts kept off every file (console only): rows the tier does not serve.
    unserved: Counter[str] = field(default_factory=Counter)


@dataclass(frozen=True)
class Stratum:
    sheet: str
    name: str
    description: str
    population: int
    sample: int
    selection: str


@dataclass
class Sample:
    sites: list[SiteItem]
    merged: list[MergedItem]
    strata: list[Stratum]
    site_frame: Frame[SiteItem]
    merged_frame: Frame[MergedItem]
    seed: int
    entitlement: str
    posture: str
    generated_at: dt.datetime
    store: str = ""
    base_url: str = ""
    largest: int = LARGEST_SITES


def _rule_index(rule: str) -> int:
    return rules.GROUPING_RULES.index(rule) if rule in rules.GROUPING_RULES else len(rules.GROUPING_RULES)


def _confidence_index(confidence: str) -> int:
    return CONFIDENCE_ORDER.index(confidence) if confidence in CONFIDENCE_ORDER else 0


def _plants(links: Iterable[Link]) -> tuple[str, ...]:
    return tuple(sorted({p for link in links if (p := link.plant_id)}, key=lambda p: (len(p), p)))


def _float(value: Any) -> float | None:
    try:
        return float(value) if value is not None and value != "" else None
    except (TypeError, ValueError):
        return None


# --------------------------------------------------------------------------- reading a store
def _stable_key(source_id: str, raw_id: str, served_id: str | None) -> str:
    """`source_id:record_id`, or `source_id:#<hash>` when the API does not print the record id (the
    hash identifies the same link in another store without printing it)."""
    if served_id is not None:
        return f"{source_id}:{served_id}"
    digest = hashlib.sha256(f"{source_id}\x1f{raw_id}".encode()).hexdigest()[:16]
    return f"{source_id}:#{digest}"


def served_link(link: ProposalSource) -> Link:
    """One readable link through the API's own member serialiser."""
    from services.api.proposal_members import member_row

    row = member_row(link)
    served_id = row["source_record_id"]
    return Link(
        source_id=row["source_id"],
        source_name=row["source_name"],
        key=_stable_key(link.source_id, link.source_record_id, served_id),
        record_id=served_id,
        name=row["name_canonical"],
        capacity_mw=_float(row["capacity_mw"]),
        technology=row["technology"],
        lifecycle_state=row["lifecycle_state"],
        role=row["role"],
        source_url=link.source_url,
    )


class PersonCheck:
    """Is any printed name a natural person's? The organisations the personal-data pass flagged, and
    the same classifier over every other name (`services.personal_names`)."""

    def __init__(self, session: Session) -> None:
        flagged = session.scalars(
            select(Organization.name_canonical).where(Organization.personal_data.is_(True))
        )
        self.flagged_names = sorted({normalise(n) for n in flagged if n and len(normalise(n)) >= 5})

    def names_person(self, *texts: str | None) -> bool:
        for raw in texts:
            if not raw:
                continue
            if is_personal_name(raw):
                return True
            norm = normalise(raw)
            if any(f" {name} " in f" {norm} " for name in self.flagged_names):
                return True
        return False


def _sponsor(view: Any) -> tuple[str | None, bool]:
    """The served sponsor's name, and whether it is a natural person."""
    org = view.sponsor  # the served view: None unless `organization_visible` admits it
    if org is None:
        return None, False
    return org.name_canonical, bool(org.personal_data)


def read_site_frame(session: Session, entitlement: str, persons: PersonCheck) -> Frame[SiteItem]:
    """Every site served at `entitlement`, as `GET /v1/sites/{id}` would serve it."""
    from services.api.interconnection_points import point_name
    from services.api.sites import served_site
    from services.api.visibility import gated_record, interconnection_point_visible, visible_source_links

    frame: Frame[SiteItem] = Frame([])
    for site in session.scalars(select(Site).where(Site.retired_at.is_(None)).order_by(Site.public_id)):
        if site.review_flag is not None:
            frame.excluded["flagged for review, not served"] += 1
            continue
        served = served_site(session, site, entitlement)
        if served is None:
            frame.unserved[f"fewer than two members served at the {entitlement} tier"] += 1
            continue
        rank_of = {m.proposal.id: m.rank for m in served.members}
        members: list[Member] = []
        personal = False
        for m in served.members:
            view = gated_record(m.proposal, entitlement)
            sponsor, sponsor_personal = _sponsor(view)
            links = tuple(served_link(link) for link in visible_source_links(m.proposal.sources, entitlement))
            point = m.proposal.interconnection_point
            poi = (
                point_name(point)
                if point is not None and interconnection_point_visible(point, entitlement)
                else None
            )
            members.append(
                Member(
                    rank=m.rank,
                    public_id=view.public_id,
                    slug=view.slug,
                    name=view.name_canonical,
                    capacity_mw=_float(view.capacity_mw),
                    technology=view.technology,
                    lifecycle_state=view.lifecycle_state,
                    jurisdiction=view.jurisdiction,
                    sponsor=sponsor,
                    relation=m.label.relation,
                    relation_rule=m.label.rule,
                    confidence=m.label.confidence,
                    grouping_rule=m.member.grouping_rule,
                    group_key=m.group_key,
                    parent_rank=rank_of.get(m.parent.id) if m.parent is not None else None,
                    connection_point=poi,
                    links=links,
                )
            )
            personal = (
                personal
                or sponsor_personal
                or persons.names_person(view.name_canonical, sponsor, poi, *(link.name for link in links))
            )
        if personal:
            frame.excluded["a printed name may be a natural person's"] += 1
            continue
        frame.items.append(SiteItem(site.public_id, members[0].name, members))
    return frame


def merged_flags(links: Sequence[Link]) -> tuple[str, ...]:
    """Every merged-record category `links` fall in, in `MERGED_CATEGORIES` order."""
    eia = [link for link in links if link.source_id == evidence.EIA_SOURCE_ID]
    plants = {p for link in eia if (p := link.plant_id)}
    per_source = Counter(link.source_id for link in links)
    markers = {evidence.phase_markers(link.name) for link in links if link.name}
    held = {
        "multi_plant": len(plants) >= 2,
        "multi_request": any(n >= 2 for sid, n in per_source.items() if sid != evidence.EIA_SOURCE_ID),
        "phase_markers": len(markers) >= 2,
        "multi_generator": len(eia) >= 2 and len(plants) <= 1,
        "cross_register": len(per_source) >= 2,
    }
    return tuple(name for name, _text in MERGED_CATEGORIES if held[name])


def read_merged_frame(
    session: Session, entitlement: str, persons: PersonCheck, site_of: Mapping[str, str]
) -> Frame[MergedItem]:
    """Every live record served at `entitlement` with two or more served source links."""
    from services.api.visibility import gated_record, proposal_visibility_filter, visible_source_links

    multi = (
        select(ProposalSource.proposal_id)
        .where(ProposalSource.active.is_(True))
        .group_by(ProposalSource.proposal_id)
        .having(func.count(ProposalSource.id) >= 2)
    )
    live = [Proposal.merged_into_id.is_(None), Proposal.id.in_(multi)]
    stored = session.scalar(select(func.count()).select_from(Proposal).where(*live)) or 0
    rows = session.scalars(
        select(Proposal)
        .where(*live, *proposal_visibility_filter(entitlement))
        .options(selectinload(Proposal.sources))
        .order_by(Proposal.public_id)
    ).all()
    frame: Frame[MergedItem] = Frame([])
    frame.unserved[f"not served at the {entitlement} tier"] = stored - len(rows)
    for p in rows:
        links = tuple(served_link(link) for link in visible_source_links(p.sources, entitlement))
        if len(links) < 2:
            frame.unserved["fewer than two links served"] += 1
            continue
        view = gated_record(p, entitlement)
        sponsor, sponsor_personal = _sponsor(view)
        if sponsor_personal or persons.names_person(
            view.name_canonical, sponsor, *(link.name for link in links)
        ):
            frame.excluded["a printed name may be a natural person's"] += 1
            continue
        frame.items.append(
            MergedItem(
                public_id=view.public_id,
                slug=view.slug,
                name=view.name_canonical,
                capacity_mw=_float(view.capacity_mw),
                technology=view.technology,
                lifecycle_state=view.lifecycle_state,
                jurisdiction=view.jurisdiction,
                sponsor=sponsor,
                links=links,
                flags=merged_flags(links),
                site_public_id=site_of.get(view.public_id),
            )
        )
    return frame


# -------------------------------------------------------------------------------- the sample
def assign_site_strata(items: Sequence[SiteItem], largest: int = LARGEST_SITES) -> None:
    """Stratum per site, the first that holds (module docstring)."""
    top = {id(s) for s in sorted(items, key=lambda s: (-len(s.members), s.key))[: max(0, largest)]}
    for s in items:
        if id(s) in top:
            s.stratum = "largest"
        elif any(m.relation == "same_site" for m in s.members):
            s.stratum = "same_site"
        else:
            s.stratum = f"{RULE_GROUP.get(s.weakest_rule, s.weakest_rule)}/{s.lowest_confidence}"


def assign_merged_strata(items: Sequence[MergedItem]) -> None:
    for m in items:
        m.stratum = m.flags[0] if m.flags else "other"


def stratum_text(sheet: str, name: str) -> str:
    if sheet == SHEET_MERGED:
        return MERGED_TEXT.get(name, "Other merged records")
    if name in SITE_STRATA_TEXT:
        return SITE_STRATA_TEXT[name]
    group, _, confidence = name.partition("/")
    return f"Grouped by {GROUP_TEXT.get(group, group)}; lowest label confidence {confidence}"


def draw(
    items: Sequence[T],
    stratum_of: Callable[[T], str],
    key_of: Callable[[T], str],
    sizes: Mapping[str, int],
    seed: int,
) -> list[T]:
    """In each stratum, the `sizes[h]` items with the smallest permanent random numbers; the result in
    one seeded random order, so strata are mixed on the sheet."""
    by: dict[str, list[T]] = {}
    for item in items:
        by.setdefault(stratum_of(item), []).append(item)
    chosen: list[T] = []
    for h, group in by.items():
        ranked = sorted(group, key=lambda it: (prn(seed, key_of(it)), key_of(it)))
        chosen.extend(ranked[: sizes.get(h, 0)])
    return sorted(chosen, key=lambda it: (prn(seed, "order\x1f" + key_of(it)), key_of(it)))


def build_sample(
    site_frame: Frame[SiteItem],
    merged_frame: Frame[MergedItem],
    *,
    seed: int = DEFAULT_SEED,
    n_sites: int = SITE_SAMPLE,
    n_merged: int = MERGED_SAMPLE,
    largest: int = LARGEST_SITES,
    entitlement: str = "public",
    posture: str = "",
    generated_at: dt.datetime | None = None,
    store: str = "",
    base_url: str = "",
) -> Sample:
    assign_site_strata(site_frame.items, largest)
    assign_merged_strata(merged_frame.items)
    site_pop = Counter(s.stratum for s in site_frame.items)
    merged_pop = Counter(m.stratum for m in merged_frame.items)
    site_sizes = allocate(site_pop, n_sites, minimum=MIN_PER_STRATUM, certainty=("largest",))
    merged_sizes = allocate(merged_pop, n_merged, minimum=MIN_PER_STRATUM)
    sites = draw(site_frame.items, lambda s: s.stratum, lambda s: s.key, site_sizes, seed)
    merged = draw(merged_frame.items, lambda m: m.stratum, lambda m: m.key, merged_sizes, seed)
    strata: list[Stratum] = []
    site_order = ["largest", "same_site", *sorted(h for h in site_pop if h not in ("largest", "same_site"))]
    for h in site_order:
        if site_pop.get(h):
            selection = "all" if site_sizes[h] >= site_pop[h] else "permanent random numbers"
            strata.append(
                Stratum(SHEET_SITES, h, stratum_text(SHEET_SITES, h), site_pop[h], site_sizes[h], selection)
            )
    for h, _text in (*MERGED_CATEGORIES, ("other", "")):
        if merged_pop.get(h):
            selection = "all" if merged_sizes[h] >= merged_pop[h] else "permanent random numbers"
            strata.append(
                Stratum(
                    SHEET_MERGED, h, stratum_text(SHEET_MERGED, h), merged_pop[h], merged_sizes[h], selection
                )
            )
    return Sample(
        sites=sites,
        merged=merged,
        strata=strata,
        site_frame=site_frame,
        merged_frame=merged_frame,
        seed=seed,
        entitlement=entitlement,
        posture=posture,
        generated_at=generated_at or dt.datetime.now(dt.UTC),
        store=store,
        base_url=base_url,
        largest=largest,
    )


# ------------------------------------------------------------------------------------- text
def mw_text(value: float | None) -> str:
    if value is None:
        return "MW not stated"
    return f"{value:,.3f}".rstrip("0").rstrip(".") + " MW"


def _tech(token: str | None) -> str:
    return (technology_label(token) or "technology not stated").lower()


def _status(token: str | None) -> str:
    return lifecycle_label(token) or "Unknown"


def relation_text(m: Member) -> str:
    template = RELATION_TEXT.get(m.relation, m.relation.replace("_", " "))
    words = template.format(parent=(m.parent_rank or 0) + 1)
    return words if m.relation == "lead" else f"{words} ({m.confidence})"


def _sources(links: Sequence[Link]) -> str:
    return "; ".join(dict.fromkeys(link.source_name for link in links)) or "no served source"


def member_line(m: Member) -> str:
    return (
        f"#{m.rank + 1} {m.name} | {_sources(m.links)} | {mw_text(m.capacity_mw)} {_tech(m.technology)}"
        f" | {_status(m.lifecycle_state)} | {relation_text(m)}"
    )


def members_text(site: SiteItem, units_shown: int = UNITS_SHOWN) -> str:
    """One line per member in served order; a plant group's units past `units_shown` in one line."""
    lines: list[str] = []
    groups: dict[str, list[Member]] = {}
    for m in site.members:
        groups.setdefault(m.group_key, []).append(m)
    for group in groups.values():
        head, units = group[0], group[1:]
        lines.append(member_line(head))
        shown = units if len(units) <= units_shown else units[: units_shown - 1]
        lines.extend("    " + member_line(u) for u in shown)
        rest = units[len(shown) :]
        if rest:
            mws = [u.capacity_mw for u in rest if u.capacity_mw is not None]
            span = f"{mw_text(min(mws))} to {mw_text(max(mws))}" if mws else "MW not stated"
            plants = ", ".join(_plants(link for u in rest for link in u.links)) or "n/a"
            statuses = ", ".join(sorted({_status(u.lifecycle_state) for u in rest}))
            lines.append(
                f"    ... and {len(rest)} more units of EIA plant {plants} (unit of #{head.rank + 1}, high):"
                f" {span}, {statuses}; all listed on '{SHEET_SITE_MEMBERS}'"
            )
    return "\n".join(lines)


def link_line(k: int, link: Link) -> str:
    rid = link.record_id if link.record_id is not None else "id not published under the licence"
    return (
        f"{k}. {link.source_name}: {link.name or 'no name'} | {mw_text(link.capacity_mw)} "
        f"{_tech(link.technology)}"
        f" | {_status(link.lifecycle_state)} | {rid}"
    )


def _one_line(item: MergedItem | Member) -> str:
    parts = [mw_text(item.capacity_mw), _tech(item.technology), _status(item.lifecycle_state)]
    if item.jurisdiction:
        parts.append(item.jurisdiction)
    return ", ".join(parts)


# ------------------------------------------------------------------------------------- rows
@dataclass(frozen=True)
class Column:
    header: str
    width: float
    kind: str = "text"  # text | wrap | int | input | input_wrap | link | url | formula
    choices: tuple[str, ...] = ()


@dataclass(frozen=True)
class LinkCell:
    """A link built from the base-URL cell (`path` relative to the site) or a full URL."""

    path: str
    label: str
    absolute: bool = False


SITE_COLUMNS: tuple[Column, ...] = (
    Column(H_ITEM, 7),
    Column("Site", 30, "link"),
    Column("Grouped because", 22, "wrap"),
    Column("Members (in site order)", 90, "wrap"),
    Column(H_SITE_VERDICT, 18, "input", SITE_VERDICTS),
    Column(H_SITE_PROBLEM, 24, "input", SITE_PROBLEMS),
    Column(H_NOTES, 40, "input_wrap"),
    Column("Lead", 30, "wrap"),
    Column("Members", 9, "int"),
    Column("Lowest label confidence", 12),
    Column("Connection points", 26, "wrap"),
    Column("Sponsors", 26, "wrap"),
    Column("EIA plant ids", 14, "wrap"),
    Column(H_STRATUM, 18),
    Column(H_POPULATION, 11, "int"),
    Column(H_SAMPLED, 10, "int"),
    Column("Weight", 9, "formula"),
    Column("Site id (this store)", 22),
    Column("Member keys (stable)", 60),
)
MERGED_COLUMNS: tuple[Column, ...] = (
    Column(H_ITEM, 7),
    Column("Record", 30, "link"),
    Column("As served", 26, "wrap"),
    Column("Source links", 88, "wrap"),
    Column(H_MERGED_VERDICT, 20, "input", MERGED_VERDICTS),
    Column(H_MERGED_KIND, 26, "input", MERGED_KINDS),
    Column(H_NOTES, 40, "input_wrap"),
    Column("Links", 7, "int"),
    Column("Registers", 9, "int"),
    Column("EIA plant ids", 14, "wrap"),
    Column("Category", 26, "wrap"),
    Column("Also matches", 26, "wrap"),
    Column("In site", 22),
    Column(H_STRATUM, 16),
    Column(H_POPULATION, 11, "int"),
    Column(H_SAMPLED, 10, "int"),
    Column("Weight", 9, "formula"),
    Column("Record id (this store)", 22),
    Column("Link keys (stable)", 60),
)
SITE_MEMBER_COLUMNS: tuple[Column, ...] = (
    Column(H_ITEM, 7),
    Column("Site", 26),
    Column("#", 5, "int"),
    Column("Member", 34, "link"),
    Column("Relation", 26),
    Column("Confidence", 11),
    Column("Relation rule", 28),
    Column("Grouped by", 30),
    Column("Sources", 28, "wrap"),
    Column("MW", 10),
    Column("Technology", 18),
    Column("Status", 16),
    Column("Jurisdiction", 11),
    Column("Sponsor", 26, "wrap"),
    Column("Connection point", 26, "wrap"),
    Column("EIA plant ids", 12),
    Column("Record id (this store)", 22),
    Column("Link keys (stable)", 60),
)
MERGED_LINK_COLUMNS: tuple[Column, ...] = (
    Column(H_ITEM, 7),
    Column("Record", 30),
    Column("#", 5, "int"),
    Column("Register", 30, "wrap"),
    Column("Name in register", 34, "wrap"),
    Column("MW", 10),
    Column("Technology", 18),
    Column("Status", 16),
    Column("Role", 10),
    Column("Source record id", 22),
    Column("EIA plant id", 12),
    Column("Register page", 18, "url"),
    Column("Stable key", 44),
)
STRATA_COLUMNS: tuple[Column, ...] = (
    Column("Sheet", 16),
    Column(H_STRATUM, 20),
    Column("Description", 60, "wrap"),
    Column("Population", 12, "int"),
    Column("Sample", 9, "int"),
    Column("Weight", 9, "formula"),
    Column("Selection", 24),
)


def site_rows(sample: Sample) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    pop = {(s.sheet, s.name): s for s in sample.strata}
    rows: list[dict[str, Any]] = []
    member_rows: list[dict[str, Any]] = []
    for k, site in enumerate(sample.sites, start=1):
        item = f"S{k:02d}"
        lead = site.members[0]
        stratum = pop[(SHEET_SITES, site.stratum)]
        rows.append(
            {
                H_ITEM: item,
                "Site": LinkCell(f"/sites/{site.public_id}", site.name),
                "Grouped because": RULE_TEXT.get(site.weakest_rule, site.weakest_rule)
                + ("" if site.weakest_rule == "eia_plant" else " (weakest link in the site)"),
                "Lead": f"{lead.name}\n{_one_line(lead)}",
                "Members (in site order)": members_text(site),
                H_SITE_VERDICT: None,
                H_SITE_PROBLEM: None,
                H_NOTES: None,
                "Members": len(site.members),
                "Lowest label confidence": site.lowest_confidence,
                "Connection points": "\n".join(
                    sorted({m.connection_point for m in site.members if m.connection_point})
                ),
                "Sponsors": "\n".join(sorted({m.sponsor for m in site.members if m.sponsor})),
                "EIA plant ids": ", ".join(site.plant_ids),
                H_STRATUM: site.stratum,
                H_POPULATION: stratum.population,
                H_SAMPLED: stratum.sample,
                "Weight": None,
                "Site id (this store)": site.public_id,
                "Member keys (stable)": "; ".join(link.key for m in site.members for link in m.links),
            }
        )
        for m in site.members:
            member_rows.append(
                {
                    H_ITEM: item,
                    "Site": site.name,
                    "#": m.rank + 1,
                    "Member": LinkCell(f"/proposals/{m.slug}", m.name),
                    "Relation": relation_text(m).rsplit(" (", 1)[0],
                    "Confidence": m.confidence,
                    "Relation rule": m.relation_rule,
                    "Grouped by": RULE_TEXT.get(m.grouping_rule, m.grouping_rule),
                    "Sources": _sources(m.links),
                    "MW": mw_text(m.capacity_mw),
                    "Technology": _tech(m.technology),
                    "Status": _status(m.lifecycle_state),
                    "Jurisdiction": m.jurisdiction or "",
                    "Sponsor": m.sponsor or "",
                    "Connection point": m.connection_point or "",
                    "EIA plant ids": ", ".join(m.plant_ids),
                    "Record id (this store)": m.public_id,
                    "Link keys (stable)": "; ".join(link.key for link in m.links),
                }
            )
    return rows, member_rows


def merged_rows(sample: Sample) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    pop = {(s.sheet, s.name): s for s in sample.strata}
    rows: list[dict[str, Any]] = []
    link_rows: list[dict[str, Any]] = []
    for k, rec in enumerate(sample.merged, start=1):
        item = f"M{k:02d}"
        stratum = pop[(SHEET_MERGED, rec.stratum)]
        served = _one_line(rec) + (f"\nSponsor: {rec.sponsor}" if rec.sponsor else "")
        rows.append(
            {
                H_ITEM: item,
                "Record": LinkCell(f"/proposals/{rec.slug}", rec.name),
                "As served": served,
                "Source links": "\n".join(link_line(i, link) for i, link in enumerate(rec.links, start=1)),
                H_MERGED_VERDICT: None,
                H_MERGED_KIND: None,
                H_NOTES: None,
                "Links": len(rec.links),
                "Registers": len({link.source_id for link in rec.links}),
                "EIA plant ids": ", ".join(rec.plant_ids),
                "Category": MERGED_TEXT.get(rec.stratum, rec.stratum),
                "Also matches": "\n".join(MERGED_TEXT[f] for f in rec.flags[1:]),
                "In site": rec.site_public_id or "",
                H_STRATUM: rec.stratum,
                H_POPULATION: stratum.population,
                H_SAMPLED: stratum.sample,
                "Weight": None,
                "Record id (this store)": rec.public_id,
                "Link keys (stable)": "; ".join(link.key for link in rec.links),
            }
        )
        for i, link in enumerate(rec.links, start=1):
            link_rows.append(
                {
                    H_ITEM: item,
                    "Record": rec.name,
                    "#": i,
                    "Register": link.source_name,
                    "Name in register": link.name or "",
                    "MW": mw_text(link.capacity_mw),
                    "Technology": _tech(link.technology),
                    "Status": _status(link.lifecycle_state),
                    "Role": link.role,
                    "Source record id": link.record_id or "not published under the licence",
                    "EIA plant id": link.plant_id or "",
                    "Register page": LinkCell(link.source_url, "Open in register", absolute=True)
                    if link.source_url
                    else None,
                    "Stable key": link.key,
                }
            )
    return rows, link_rows


def strata_rows(sample: Sample) -> list[dict[str, Any]]:
    return [
        {
            "Sheet": s.sheet,
            H_STRATUM: s.name,
            "Description": s.description,
            "Population": s.population,
            "Sample": s.sample,
            "Weight": None,
            "Selection": s.selection,
        }
        for s in sample.strata
    ]


# ----------------------------------------------------------------------------------- files
def _csv_cells(column: Column, value: Any) -> list[Any]:
    """A link column is two CSV columns: its label, then `<header> link` (a path relative to the
    base URL, or the register's own URL)."""
    if column.kind in ("link", "url"):
        if isinstance(value, LinkCell):
            return [value.label, value.path]
        return ["", ""]
    return ["" if value is None else value]


def write_csv(path: pathlib.Path, columns: Sequence[Column], rows: Sequence[Mapping[str, Any]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        header: list[str] = []
        for c in columns:
            header += [c.header, f"{c.header} link"] if c.kind in ("link", "url") else [c.header]
        writer.writerow(header)
        for row in rows:
            writer.writerow([cell for c in columns for cell in _csv_cells(c, row.get(c.header))])


def _letter(index: int) -> str:
    from openpyxl.utils import get_column_letter

    return str(get_column_letter(index))


def _quote(text_: str) -> str:
    return text_.replace('"', '""')


def _row_height(row: Mapping[str, Any], columns: Sequence[Column]) -> float:
    lines = 1
    for c in columns:
        value = row.get(c.header)
        if c.kind in ("wrap", "input_wrap") and isinstance(value, str) and value:
            per_line = max(8, int(c.width * 1.15))
            lines = max(lines, sum(max(1, -(-len(line) // per_line)) for line in value.split("\n")))
    return float(min(409, max(15, 13 * lines + 4)))


class _Styles:
    def __init__(self) -> None:
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        self.font = Font(name="Arial", size=10)
        self.bold = Font(name="Arial", size=10, bold=True)
        self.header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        self.title = Font(name="Arial", size=14, bold=True)
        self.heading = Font(name="Arial", size=11, bold=True, color="1F3864")
        self.link = Font(name="Arial", size=10, color="0563C1", underline="single")
        self.header_fill = PatternFill("solid", fgColor="1F3864")
        self.input_fill = PatternFill("solid", fgColor="FFF2CC")
        self.bookkeeping_fill = PatternFill("solid", fgColor="F2F2F2")
        self.wrap = Alignment(wrap_text=True, vertical="top")
        self.top = Alignment(vertical="top")
        self.header_align = Alignment(wrap_text=True, vertical="center")
        thin = Side(style="thin", color="BFBFBF")
        self.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def write_sheet(
    ws: Any,
    columns: Sequence[Column],
    rows: Sequence[Mapping[str, Any]],
    styles: _Styles,
    *,
    weight_from: tuple[str, str] | None = None,
    freeze: str = "B2",
) -> None:
    """Header row, frozen panes, widths, wrapped text, dropdowns on input columns, HYPERLINK formulas
    from the `BaseURL` cell, and `Weight` = population / sample as a formula."""
    from openpyxl.worksheet.datavalidation import DataValidation

    index = {c.header: i for i, c in enumerate(columns, start=1)}
    for i, c in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=i, value=c.header)
        cell.font, cell.fill, cell.alignment, cell.border = (
            styles.header_font,
            styles.header_fill,
            styles.header_align,
            styles.border,
        )
        ws.column_dimensions[_letter(i)].width = c.width
    ws.row_dimensions[1].height = 42
    for r, row in enumerate(rows, start=2):
        for i, c in enumerate(columns, start=1):
            value = row.get(c.header)
            cell = ws.cell(row=r, column=i)
            cell.font, cell.border = styles.font, styles.border
            cell.alignment = styles.wrap if c.kind in ("wrap", "input_wrap", "link") else styles.top
            if c.kind in ("input", "input_wrap"):
                cell.fill = styles.input_fill
            if c.kind == "formula" and weight_from is not None:
                pop, sam = (_letter(index[h]) for h in weight_from)
                cell.value = f'=IF({sam}{r}>0,{pop}{r}/{sam}{r},"")'
                cell.number_format = "0.0"
                cell.fill = styles.bookkeeping_fill
            elif isinstance(value, LinkCell):
                target = f'"{_quote(value.path)}"' if value.absolute else f'BaseURL&"{_quote(value.path)}"'
                cell.value = f'=HYPERLINK({target},"{_quote(value.label[:200])}")'
                cell.font = styles.link
            else:
                cell.value = value
        ws.row_dimensions[r].height = _row_height(row, columns)
    last = max(2, len(rows) + 1)
    for i, c in enumerate(columns, start=1):
        if c.choices:
            dv = DataValidation(
                type="list",
                formula1='"' + ",".join(c.choices) + '"',
                allow_blank=True,
                showErrorMessage=True,
                errorTitle="Pick from the list",
                error="Choose one of: " + ", ".join(c.choices),
            )
            ws.add_data_validation(dv)
            dv.add(f"{_letter(i)}2:{_letter(i)}{last}")
    ws.freeze_panes = freeze
    ws.page_setup.orientation = "landscape"
    ws.print_title_rows = "1:1"
    ws.auto_filter.ref = f"A1:{_letter(len(columns))}{last}"


def instructions_rows(sample: Sample, site_n: int, merged_n: int) -> list[tuple[str, str]]:
    """(label, text) rows of the Instructions sheet below its title and base-URL cell. A label
    starting with '#' is a section heading."""
    clean, over, sure = decisive_counts(site_n) if site_n else (-1, 1, 1)
    zero_hi = wilson(0, site_n)[1] if site_n else 1.0
    five_lo, five_hi = wilson(min(5, site_n), site_n) if site_n else (0.0, 1.0)
    largest = sum(1 for s in sample.sites if s.stratum == "largest")
    return [
        ("#What this is for", ""),
        (
            "Purpose",
            "You are checking a sample of how the platform groups filings. Two questions decide two things "
            "for "
            "the beta: (1) do the records in a site belong to one project or site? If more than about 1 in "
            "10 "
            "sites group records that are not the same project, sites are switched off for the beta "
            "(SITES_ENABLED=0). (2) Is each merged record one real project? If many are several projects, "
            "merged records need splitting before the site can claim 'one record per project'.",
        ),
        (
            "Time",
            f"About 2 to 3 hours in all: {site_n} sites and {merged_n} merged records. Most rows take 1 to 2 "
            f"minutes; the {largest} largest sites can take 5 to 10 minutes each (their unit lists are "
            "long). "
            "You can stop and resume: blank rows are simply not counted.",
        ),
        (
            "What to fill in",
            "Only the yellow cells on the Sites and Merged records sheets: the two dropdown columns and "
            "Notes. "
            "Everything else is the platform's own data as a public visitor sees it. 'Site members' and "
            "'Merged links' list every member and source link in full, for reference.",
        ),
        ("#How to check a site (Sites sheet)", ""),
        (
            "Read the row",
            "'Members' lists the records in the site in display order: #1 is the lead (the largest active "
            "filing). Each line gives the name, register, MW, technology, status, and how the platform "
            "labels "
            "it relative to the lead, with its confidence. Units of one EIA plant are indented under the "
            "plant's "
            "first record; long unit lists are summarised. Open the site link to see the page itself.",
        ),
        (
            H_SITE_VERDICT,
            "Yes: every record belongs to one project or one physical site (phases, units, co-located "
            "technologies, re-filings and expansions of one site all count as belonging). No: at least one "
            "record does not belong with the others (a different developer's project, a competing request at "
            "the same substation, an unrelated building). Unsure: you cannot tell from public information.",
        ),
        (
            H_SITE_PROBLEM,
            "When you answer No, pick the main problem. 'Unrelated record grouped' is the one the 1-in-10 "
            "rule "
            "counts (as do 'Other' and a No with nothing picked). 'Missing record' (a record that should be "
            "in "
            "the site is not), 'Wrong lead' and 'Wrong label' are recorded as quality issues of a site whose "
            "members do belong together; you may also pick them on a Yes row. Use Notes for anything else.",
        ),
        ("#How to check a merged record (Merged records sheet)", ""),
        (
            "Read the row",
            "'Source links' lists each register row the platform merged into this one record: register, the "
            "name the register uses, MW, technology, status and the register's own id. The record link opens "
            "the record page, which shows the same members.",
        ),
        (
            H_MERGED_VERDICT,
            "Yes: all the source rows describe one real project. No (several projects): they describe two or "
            "more projects, phases or plants that should be separate records (they may still sit in one "
            "site). "
            "Unsure: you cannot tell.",
        ),
        (
            H_MERGED_KIND,
            "When the answer is No, what the rows are: phases of one development; separate units or plants "
            "at "
            "one site; a re-filing or replacement of the same project; unrelated projects; or other.",
        ),
        ("#How the result is used", ""),
        (
            "Scoring",
            "Run: python -m services.sites.handcheck score <this file>. It reports the share of sites judged "
            "No (grouping errors) with a 95 % Wilson interval, the same share weighted back to every served "
            "site (the sample over-represents the riskier groupings, so the weighted figure is the one that "
            "answers 'how many sites are wrong'), and the merged-record share judged 'several projects', by "
            "category. Unsure and blank rows are left out of the shares and reported.",
        ),
        (
            "The 1-in-10 rule",
            "Sites are switched off for the beta when the weighted estimate of sites grouping unrelated "
            "records "
            "is above 10 %. The scorer states both the point estimate and the interval's lower bound: above "
            "10 "
            "% on the point estimate meets the rule as stated; a lower bound above 10 % means the sample "
            "rules out an acceptable rate.",
        ),
        (
            "What a result will mean",
            f"With {site_n} sites judged (unweighted, before weighting): 0 wrong gives an interval of 0 to "
            f"{zero_hi:.1%}, so the rule is clearly met; up to {clean} wrong keeps the whole interval at or "
            f"under 10 %; {over} or more wrong puts the point estimate above 10 %; {sure} or more wrong puts "
            f"even the lower bound above 10 %. 5 wrong (exactly 1 in 10) gives {five_lo:.1%} to "
            f"{five_hi:.1%}: "
            f"{site_n} items cannot separate 8 % from 12 %, only clean results from clearly bad ones.",
        ),
        ("#Labels you will see", ""),
        ("LEAD", "The site's headline record: the largest filing that is not withdrawn, ties to the latest."),
        ("unit of #n", "Another generator of the same EIA plant as member #n (high confidence)."),
        ("phase of the lead", "Same project name, different phase or unit number."),
        ("co-located with the lead", "A different technology at the same site (e.g. gas beside nuclear)."),
        ("expansion of / expanded by", "An addition to an operating plant, or a stated MW increase."),
        ("superseded by / re-filing of", "A withdrawn filing that the lead replaced, or that came after it."),
        (
            "same site, relation unclear",
            "Grouped by a shared identifier, but no rule says how they relate (low).",
        ),
        (
            "Confidence",
            "high / medium / low is the rule's own strength, not a measured probability: this check is "
            "the first labelled sample.",
        ),
        ("#Example of a filled row (illustration only, not part of the sample)", ""),
        (
            "Sites example",
            "S00 | Example Ridge Solar | members: #1 Example Ridge Solar 200 MW; #2 Example Ridge Solar II "
            "100 MW (phase of the lead, high); #3 Other Co Storage 50 MW (same site, relation unclear, low) "
            "| "
            "All these belong to one project or site? = No | If no, what is wrong = Unrelated record "
            "grouped | "
            "Notes = '#3 is a different developer's battery at the same substation'",
        ),
        (
            "Merged example",
            "M00 | Example Flats | 1. CAISO queue: Example Flats I; 2. EIA-860M: Example Flats II | One real "
            "project? = No (several projects) | What kind = Phases of one development | Notes = 'two phases, "
            "separate EIA plants'",
        ),
        ("#About this file", ""),
        (
            "Produced from",
            f"{sample.store or 'a store'}; generated {sample.generated_at:%Y-%m-%d %H:%M} UTC; seed "
            f"{sample.seed}; "
            f"tier '{sample.entitlement}'; platform posture '{sample.posture or 'environment default'}'; "
            "site "
            f"rules {rules.RULE_VERSION}.",
        ),
        (
            "Stable identifiers",
            "Record and site ids are minted when a store is loaded, so another store (the beta) has "
            "different "
            "ids and links. Each row also carries stable keys (register id plus the register's record id, a "
            "hash where the licence does not let the platform print the id; EIA plant ids). Regenerate "
            "against "
            "the store you will browse: python -m services.sites.handcheck generate --database-url <url> "
            "--out "
            f"<dir> --seed {sample.seed}. The same seed draws the same items wherever the stores agree.",
        ),
    ]


def write_instructions(ws: Any, sample: Sample, styles: _Styles, progress: list[tuple[str, str, str]]) -> str:
    """The Instructions sheet; returns the base-URL cell's absolute reference."""
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 120
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.cell(row=1, column=1, value="Hand check: sites and merged records").font = styles.title
    ws.cell(
        row=2, column=1, value="Yellow cells are yours to fill in; everything else is generated."
    ).font = styles.font
    ws.cell(row=4, column=1, value="Base URL").font = styles.bold
    base = ws.cell(row=4, column=2, value=sample.base_url or None)
    base.fill, base.font, base.border = styles.input_fill, styles.font, styles.border
    note = ws.cell(
        row=5,
        column=2,
        value="Type the site's address here, without a trailing slash (e.g. https://beta.example.org). Every "
        "link in the workbook starts with it; left empty, links are relative paths (/sites/..., "
        "/proposals/...).",
    )
    note.font, note.alignment = styles.font, styles.wrap
    r = 7
    ws.cell(row=r, column=1, value="Progress (updates as you fill in)").font = styles.heading
    r += 1
    for label, formula, fmt in progress:
        head = ws.cell(row=r, column=1, value=label)
        head.font, head.alignment = styles.font, styles.wrap
        cell = ws.cell(row=r, column=2, value=formula)
        cell.font, cell.alignment = styles.font, styles.top
        cell.number_format = fmt
        r += 1
    r += 1
    site_n = len(sample.sites)
    for label, body in instructions_rows(sample, site_n, len(sample.merged)):
        if label.startswith("#"):
            r += 1
            ws.cell(row=r, column=1, value=label[1:]).font = styles.heading
        else:
            a = ws.cell(row=r, column=1, value=label)
            a.font, a.alignment = styles.bold, styles.wrap
            b = ws.cell(row=r, column=2, value=body)
            b.font, b.alignment = styles.font, styles.wrap
            ws.row_dimensions[r].height = float(max(15, 13 * (-(-len(body) // 135)) + 4))
        r += 1
    ws.freeze_panes = "A2"
    return f"{SHEET_INSTRUCTIONS}!$B$4"


def _progress(site_rows_: int, merged_rows_: int) -> list[tuple[str, str, str]]:
    sv = _letter(next(i for i, c in enumerate(SITE_COLUMNS, 1) if c.header == H_SITE_VERDICT))
    mv = _letter(next(i for i, c in enumerate(MERGED_COLUMNS, 1) if c.header == H_MERGED_VERDICT))
    # Never reach the header row: an empty sheet still counts rows 2..2 (as `write_sheet` does).
    s_rng = f"{SHEET_SITES}!{sv}2:{sv}{max(2, site_rows_ + 1)}"
    m_rng = f"'{SHEET_MERGED}'!{mv}2:{mv}{max(2, merged_rows_ + 1)}"
    return [
        ("Sites answered", f'=COUNTA({s_rng})&" of {site_rows_}"', "@"),
        ("Sites answered No", f'=COUNTIF({s_rng},"No")', "0"),
        (
            "Sites: share of Yes/No answers that are No (unweighted)",
            f'=IFERROR(COUNTIF({s_rng},"No")/(COUNTIF({s_rng},"Yes")+COUNTIF({s_rng},"No")),"")',
            "0.0%",
        ),
        ("Merged records answered", f'=COUNTA({m_rng})&" of {merged_rows_}"', "@"),
        (
            "Merged records answered 'No (several projects)'",
            f'=COUNTIF({m_rng},"No (several projects)")',
            "0",
        ),
        (
            "Merged: share of Yes/No answers that are No (unweighted)",
            f'=IFERROR(COUNTIF({m_rng},"No (several projects)")/(COUNTIF({m_rng},"Yes")'
            f'+COUNTIF({m_rng},"No (several projects)")),"")',
            "0.0%",
        ),
    ]


def write_workbook(path: pathlib.Path, sample: Sample) -> None:
    from openpyxl import Workbook
    from openpyxl.workbook.defined_name import DefinedName

    styles = _Styles()
    srows, mrows_members = site_rows(sample)
    mrows, mlinks = merged_rows(sample)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_INSTRUCTIONS
    base_ref = write_instructions(ws, sample, styles, _progress(len(srows), len(mrows)))
    wb.defined_names["BaseURL"] = DefinedName("BaseURL", attr_text=base_ref)
    weight = (H_POPULATION, H_SAMPLED)
    write_sheet(wb.create_sheet(SHEET_SITES), SITE_COLUMNS, srows, styles, weight_from=weight, freeze="C2")
    write_sheet(wb.create_sheet(SHEET_MERGED), MERGED_COLUMNS, mrows, styles, weight_from=weight, freeze="C2")
    write_sheet(wb.create_sheet(SHEET_SITE_MEMBERS), SITE_MEMBER_COLUMNS, mrows_members, styles, freeze="E2")
    write_sheet(wb.create_sheet(SHEET_MERGED_LINKS), MERGED_LINK_COLUMNS, mlinks, styles, freeze="D2")
    strata_ws = wb.create_sheet(SHEET_STRATA)
    write_sheet(
        strata_ws,
        STRATA_COLUMNS,
        strata_rows(sample),
        styles,
        weight_from=("Population", "Sample"),
        freeze="A2",
    )
    r = len(sample.strata) + 3
    for label, body in strata_notes(sample):
        a = strata_ws.cell(row=r, column=1, value=label)
        a.font = styles.bold
        b = strata_ws.cell(row=r, column=3, value=body)
        b.font, b.alignment = styles.font, styles.wrap
        strata_ws.row_dimensions[r].height = float(max(15, 13 * (-(-len(body) // 70)) + 4))
        r += 1
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


def strata_notes(sample: Sample) -> list[tuple[str, str]]:
    sites_total = len(sample.site_frame.items)
    merged_total = len(sample.merged_frame.items)
    notes = [
        (
            "Sites frame",
            f"{sites_total} sites served at the '{sample.entitlement}' tier with two or more served members.",
        ),
        (
            "Merged frame",
            f"{merged_total} live records served at the '{sample.entitlement}' tier with two or more served "
            "source links.",
        ),
    ]
    for label, frame in (("Sites", sample.site_frame), ("Merged records", sample.merged_frame)):
        for reason, n in sorted(frame.excluded.items()):
            notes.append((f"{label} left out", f"{n}: {reason}."))
    notes += [
        (
            "Allocation",
            f"The {sample.largest} largest sites are all checked (weight 1). The rest of each budget is "
            "shared in "
            "proportion to the square root of each stratum's size, at least "
            f"{MIN_PER_STRATUM} per stratum: every grouping rule and confidence level is seen, and the large "
            "strata still dominate. Weight = population / sample; the scorer reweights by the items actually "
            "judged.",
        ),
        (
            "Selection",
            f"Within a stratum, the items with the smallest permanent random numbers: SHA-256 of seed "
            f"{sample.seed} and the item's stable key (its smallest source-link key). Same seed and store: "
            "same "
            "items; another store: the same items wherever the two agree.",
        ),
    ]
    return notes


def write_outputs(out: pathlib.Path, sample: Sample, *, command: str = "", note: str = "") -> dict[str, Any]:
    """The workbook, its CSV twins, `manifest.json` and README.md under `out`; returns the manifest."""
    out.mkdir(parents=True, exist_ok=True)
    srows, member_rows_ = site_rows(sample)
    mrows, link_rows = merged_rows(sample)
    write_workbook(out / WORKBOOK, sample)
    write_csv(out / "sites.csv", SITE_COLUMNS, srows)
    write_csv(out / "merged_records.csv", MERGED_COLUMNS, mrows)
    write_csv(out / "site_members.csv", SITE_MEMBER_COLUMNS, member_rows_)
    write_csv(out / "merged_links.csv", MERGED_LINK_COLUMNS, link_rows)
    strata = strata_rows(sample)
    for row in strata:
        row["Weight"] = round(row["Population"] / row["Sample"], 4) if row["Sample"] else ""
    write_csv(out / "strata.csv", STRATA_COLUMNS, strata)
    manifest: dict[str, Any] = {
        "generated_at": sample.generated_at.isoformat(timespec="seconds"),
        "store": sample.store,
        "seed": sample.seed,
        "tier": sample.entitlement,
        "platform_posture": sample.posture,
        "site_rule_version": rules.RULE_VERSION,
        "command": command,
        "sites": {
            "frame": len(sample.site_frame.items),
            "sample": len(sample.sites),
            "excluded": dict(sample.site_frame.excluded),
        },
        "merged_records": {
            "frame": len(sample.merged_frame.items),
            "sample": len(sample.merged),
            "excluded": dict(sample.merged_frame.excluded),
        },
        "strata": [
            {
                "sheet": s.sheet,
                "stratum": s.name,
                "population": s.population,
                "sample": s.sample,
                "selection": s.selection,
            }
            for s in sample.strata
        ],
    }
    files = sorted(p for p in out.iterdir() if p.is_file() and p.name not in ("manifest.json", "README.md"))
    manifest["files"] = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in files}
    (out / "manifest.json").write_text(json.dumps(manifest, indent=1) + "\n", encoding="utf-8")
    (out / "README.md").write_text(readme(sample, command, note), encoding="utf-8")
    return manifest


def readme(sample: Sample, command: str, note: str = "") -> str:
    rows = "\n".join(
        f"| {s.sheet} | `{s.name}` | {s.description} | {s.population} | {s.sample} | {s.selection} |"
        for s in sample.strata
    )
    lead = f"{note.strip()}\n\n" if note.strip() else ""
    return f"""# Hand check worksheet ({sample.generated_at:%Y-%m-%d})

{lead}Produced by `{command or "python -m services.sites.handcheck generate"}`
from {sample.store or "a store"}, read only, at {sample.generated_at:%Y-%m-%d %H:%M} UTC
(seed {sample.seed}, tier `{sample.entitlement}`, posture `{sample.posture or "environment default"}`,
site rules `{rules.RULE_VERSION}`).

**Record and site ids and links belong to this store.** Ids are minted at load, so for working
links regenerate against the store you will browse (the beta) with the same seed, and put its
address in the base-URL cell on the Instructions sheet. Each row also carries stable keys
(`source_id:source_record_id`, a hash where the licence does not let the platform print the id;
EIA plant ids), so items can be matched across stores.

Files: `{WORKBOOK}` (fill in the yellow cells); CSV twins of each data sheet (`sites.csv`,
`merged_records.csv`, `site_members.csv`, `merged_links.csv`, `strata.csv`); `manifest.json`
(counts, strata, file hashes). Score with `python -m services.sites.handcheck score {WORKBOOK}`.

Sites: {len(sample.sites)} of {len(sample.site_frame.items)} served. Merged records:
{len(sample.merged)} of {len(sample.merged_frame.items)} served with two or more source links.

| Sheet | Stratum | What | Population | Sample | Selection |
|---|---|---|---|---|---|
{rows}
"""


# ------------------------------------------------------------------------------ generation
def _read_only(session: Session) -> None:
    dialect = session.get_bind().dialect.name
    if dialect == "postgresql":
        session.execute(text("SET TRANSACTION READ ONLY"))
    elif dialect == "sqlite":
        session.execute(text("PRAGMA query_only = ON"))


def read_frames(session: Session, entitlement: str) -> tuple[Frame[SiteItem], Frame[MergedItem]]:
    persons = PersonCheck(session)
    site_frame = read_site_frame(session, entitlement, persons)
    site_of = {m.public_id: s.public_id for s in site_frame.items for m in s.members}
    merged_frame = read_merged_frame(session, entitlement, persons, site_of)
    return site_frame, merged_frame


def generate(
    session: Session,
    out: pathlib.Path,
    *,
    seed: int = DEFAULT_SEED,
    entitlement: str = "public",
    n_sites: int = SITE_SAMPLE,
    n_merged: int = MERGED_SAMPLE,
    largest: int = LARGEST_SITES,
    posture: str = "",
    store: str = "",
    base_url: str = "",
    command: str = "",
    note: str = "",
    now: dt.datetime | None = None,
) -> Sample:
    """Read the frames from `session` (the caller makes it read-only and rolls back), draw the sample
    and write every file under `out`."""
    site_frame, merged_frame = read_frames(session, entitlement)
    sample = build_sample(
        site_frame,
        merged_frame,
        seed=seed,
        n_sites=n_sites,
        n_merged=n_merged,
        largest=largest,
        entitlement=entitlement,
        posture=posture,
        generated_at=now,
        store=store,
        base_url=base_url,
    )
    write_outputs(out, sample, command=command, note=note)
    return sample


def describe(sample: Sample) -> str:
    lines = [
        f"sites: {len(sample.sites)} drawn from {len(sample.site_frame.items)} served",
        f"merged records: {len(sample.merged)} drawn from {len(sample.merged_frame.items)} served",
    ]
    for label, frame in (("sites", sample.site_frame), ("merged", sample.merged_frame)):
        for reason, n in sorted({**frame.excluded, **frame.unserved}.items()):
            lines.append(f"  {label} left out: {n} ({reason})")
    for s in sample.strata:
        lines.append(
            f"  {s.sheet:<15} {s.name:<22} population {s.population:>5}  sample {s.sample:>3}  {s.selection}"
        )
    return "\n".join(lines)


# ---------------------------------------------------------------------------------- scoring
def _norm(value: Any) -> str:
    return " ".join(str(value).split()).lower() if value is not None else ""


def _int(value: Any) -> int:
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return 0


def read_filled(path: pathlib.Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """The Sites and Merged records rows of a filled workbook, or of `sites.csv` and
    `merged_records.csv` in a directory."""
    if path.is_dir():
        out: list[list[dict[str, Any]]] = []
        for name in ("sites.csv", "merged_records.csv"):
            with (path / name).open(newline="", encoding="utf-8") as fh:
                out.append([dict(row) for row in csv.DictReader(fh)])
        return out[0], out[1]
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True)
    try:
        sheets = []
        for name in (SHEET_SITES, SHEET_MERGED):
            values = list(wb[name].iter_rows(values_only=True))
            header = [str(h) if h is not None else "" for h in values[0]] if values else []
            sheets.append([dict(zip(header, row, strict=False)) for row in values[1:]])
        return sheets[0], sheets[1]
    finally:
        wb.close()


@dataclass
class Score:
    sites: Estimate
    sites_any_no: Estimate
    sites_unsure_as_wrong: Estimate
    merged: Estimate
    site_problems: Counter[str]
    site_issues_on_yes: Counter[str]
    merged_kinds: Counter[str]
    unreadable: list[str]

    def kill_switch(self) -> dict[str, Any]:
        point = self.sites.weighted_rate
        low, high = self.sites.weighted_interval
        raw_low, raw_high = self.sites.interval
        return {
            "threshold": KILL_SWITCH_RATE,
            "weighted_point": point,
            "weighted_lower": low,
            "weighted_upper": high,
            "crossed_on_point": point is not None and point > KILL_SWITCH_RATE,
            "crossed_on_lower_bound": point is not None and low > KILL_SWITCH_RATE,
            "sample_point": self.sites.rate,
            "sample_lower": raw_low,
            "sample_upper": raw_high,
            "sample_crossed_on_point": self.sites.rate is not None and self.sites.rate > KILL_SWITCH_RATE,
            "sample_crossed_on_lower_bound": self.sites.rate is not None and raw_low > KILL_SWITCH_RATE,
        }

    def as_dict(self) -> dict[str, Any]:
        return {
            "sites": self.sites.as_dict(),
            "sites_any_no": self.sites_any_no.as_dict(),
            "sites_unsure_as_wrong": self.sites_unsure_as_wrong.as_dict(),
            "kill_switch": self.kill_switch(),
            "site_problems": dict(self.site_problems),
            "site_issues_on_yes": dict(self.site_issues_on_yes),
            "merged": self.merged.as_dict(),
            "merged_kinds": dict(self.merged_kinds),
            "unreadable": self.unreadable,
        }


def score_rows(site_rows_: Sequence[Mapping[str, Any]], merged_rows_: Sequence[Mapping[str, Any]]) -> Score:
    """Grouping error = a site answered No whose problem is 'Unrelated record grouped', 'Other' or
    none; merged error = 'No (several projects)'. Unsure and blank are left out (`handcheck_stats`)."""
    unreadable: list[str] = []
    kill: list[Judged] = []
    any_no: list[Judged] = []
    unsure_wrong: list[Judged] = []
    problems: Counter[str] = Counter()
    issues_on_yes: Counter[str] = Counter()
    site_ok = {v.lower() for v in SITE_VERDICTS}
    for row in site_rows_:
        item = str(row.get(H_ITEM) or "").strip()
        if not item.startswith("S"):
            continue
        stratum, pop, sam = (
            str(row.get(H_STRATUM) or ""),
            _int(row.get(H_POPULATION)),
            _int(row.get(H_SAMPLED)),
        )
        verdict, problem = _norm(row.get(H_SITE_VERDICT)), _norm(row.get(H_SITE_PROBLEM))
        if verdict and verdict not in site_ok:
            unreadable.append(f"{item}: {row.get(H_SITE_VERDICT)!r}")
            verdict = ""
        if verdict == "no":
            problems[problem or "(none given)"] += 1
        elif verdict == "yes" and problem:
            issues_on_yes[problem] += 1
        unsure = verdict == "unsure"
        judged = verdict in ("yes", "no")
        grouping_error = verdict == "no" and problem in GROUPING_PROBLEMS
        kill.append(Judged(stratum, pop, sam, grouping_error if judged else None, unsure))
        any_no.append(Judged(stratum, pop, sam, (verdict == "no") if judged else None, unsure))
        unsure_wrong.append(
            Judged(stratum, pop, sam, (grouping_error or unsure) if (judged or unsure) else None, False)
        )
    merged: list[Judged] = []
    kinds: Counter[str] = Counter()
    merged_ok = {v.lower() for v in MERGED_VERDICTS}
    for row in merged_rows_:
        item = str(row.get(H_ITEM) or "").strip()
        if not item.startswith("M"):
            continue
        verdict = _norm(row.get(H_MERGED_VERDICT))
        if verdict and verdict not in merged_ok:
            unreadable.append(f"{item}: {row.get(H_MERGED_VERDICT)!r}")
            verdict = ""
        several = verdict.startswith("no")
        if several:
            kinds[str(row.get(H_MERGED_KIND) or "(none given)").strip()] += 1
        judged = verdict == "yes" or several
        merged.append(
            Judged(
                str(row.get(H_STRATUM) or ""),
                _int(row.get(H_POPULATION)),
                _int(row.get(H_SAMPLED)),
                several if judged else None,
                verdict == "unsure",
            )
        )
    return Score(
        sites=estimate(kill),
        sites_any_no=estimate(any_no),
        sites_unsure_as_wrong=estimate(unsure_wrong),
        merged=estimate(merged),
        site_problems=problems,
        site_issues_on_yes=issues_on_yes,
        merged_kinds=kinds,
        unreadable=unreadable,
    )


def _pct(value: float | None) -> str:
    return "n/a" if value is None else f"{value:.1%}"


def _ci(interval: tuple[float, float]) -> str:
    return f"{interval[0]:.1%} to {interval[1]:.1%}"


def _counts(e: Estimate) -> str:
    return f"{e.sampled} in the sample: {e.judged} judged, {e.unsure} unsure, {e.blank} blank"


def render_score(score: Score) -> str:
    s, m, k = score.sites, score.merged, score.kill_switch()
    lines = ["SITES", f"  {_counts(s)}"]
    lines.append(f"  grouping errors: {s.errors} (No with 'Unrelated record grouped', 'Other' or no problem)")
    lines.append(f"  sample share wrong: {_pct(s.rate)} (Wilson 95 % CI {_ci(s.interval)})")
    lines.append(
        f"  weighted to all served sites: {_pct(s.weighted_rate)} (95 % CI {_ci(s.weighted_interval)};"
        f" effective n {s.effective_n:.1f}; strata judged cover {s.covered:.0%} of {s.population} sites)"
    )
    verdict = (
        "CROSSED: switch sites off for the beta (SITES_ENABLED=0)"
        if k["crossed_on_point"]
        else "not crossed: sites stay on"
    )
    lines.append(f"  1-in-10 rule on the weighted point estimate {_pct(k['weighted_point'])}: {verdict}")
    sure = "yes" if k["crossed_on_lower_bound"] else "no"
    lines.append(f"  ... and on the interval's lower bound {_pct(k['weighted_lower'])}: crossed = {sure}")
    lines.append(
        f"  unweighted, for reference: point {_pct(k['sample_point'])} crossed = "
        f"{'yes' if k['sample_crossed_on_point'] else 'no'}; lower bound {_pct(k['sample_lower'])} crossed = "
        f"{'yes' if k['sample_crossed_on_lower_bound'] else 'no'}"
    )
    lines.append(
        f"  any No (incl. missing record, wrong lead, wrong label): {_pct(score.sites_any_no.weighted_rate)}"
        f" weighted ({_ci(score.sites_any_no.weighted_interval)})"
    )
    lines.append(
        f"  unsure counted as wrong: {_pct(score.sites_unsure_as_wrong.weighted_rate)} weighted"
        f" ({_ci(score.sites_unsure_as_wrong.weighted_interval)})"
    )
    if score.site_problems:
        lines.append(
            "  problems on No: " + ", ".join(f"{p} {n}" for p, n in score.site_problems.most_common())
        )
    if score.site_issues_on_yes:
        lines.append(
            "  issues on Yes: " + ", ".join(f"{p} {n}" for p, n in score.site_issues_on_yes.most_common())
        )
    lines += _strata_lines(s)
    lines += ["", "MERGED RECORDS", f"  {_counts(m)}"]
    lines.append(
        f"  several projects: {m.errors}; sample share {_pct(m.rate)} (Wilson 95 % CI {_ci(m.interval)})"
    )
    lines.append(
        f"  weighted to all served merged records: {_pct(m.weighted_rate)} (95 % CI "
        f"{_ci(m.weighted_interval)};"
        f" effective n {m.effective_n:.1f}; strata judged cover {m.covered:.0%} of {m.population})"
    )
    if score.merged_kinds:
        lines.append("  what kind: " + ", ".join(f"{p} {n}" for p, n in score.merged_kinds.most_common()))
    lines += _strata_lines(m)
    if score.unreadable:
        lines += ["", "Unreadable verdicts (counted as blank): " + "; ".join(score.unreadable)]
    return "\n".join(lines)


def _strata_lines(e: Estimate) -> list[str]:
    out = ["  by stratum (population, sampled, judged, errors, share, Wilson 95 % CI):"]
    for st in e.strata:
        out.append(
            f"    {st.stratum:<22} {st.population:>5} {st.sampled:>4} {st.judged:>4} {st.errors:>4}"
            f"  {_pct(st.rate):>6}  {_ci(st.interval) if st.judged else 'n/a'}"
        )
    return out


# -------------------------------------------------------------------------------------- CLI
def _redacted(url: str | None) -> str:
    from sqlalchemy.engine import make_url

    raw = url or os.environ.get("DATABASE_URL") or ""
    try:
        return make_url(raw).render_as_string(hide_password=True) if raw else "the DATABASE_URL store"
    except Exception:  # a description only; never print the raw value
        return "the given store"


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="python -m services.sites.handcheck", description=__doc__.split("\n")[0]
    )
    sub = parser.add_subparsers(dest="cmd", required=True)
    gen = sub.add_parser("generate", help="draw the sample from a store (read only) and write the worksheet")
    gen.add_argument("--database-url", default=None, help="defaults to DATABASE_URL")
    gen.add_argument("--out", type=pathlib.Path, default=None, help="default data/eval/handcheck/<today>")
    gen.add_argument("--seed", type=int, default=DEFAULT_SEED)
    gen.add_argument("--sites", type=int, default=SITE_SAMPLE)
    gen.add_argument("--merged", type=int, default=MERGED_SAMPLE)
    gen.add_argument("--largest", type=int, default=LARGEST_SITES, help="largest sites checked in full")
    gen.add_argument("--tier", default="public", choices=("public", "pro", "api"))
    gen.add_argument(
        "--posture",
        choices=("commercial", "noncommercial"),
        default=None,
        help="the PLATFORM_POSTURE the site is served under (the beta: noncommercial); default: environment",
    )
    gen.add_argument("--base-url", default="", help="prefills the workbook's base-URL cell")
    gen.add_argument("--note", default="", help="a paragraph put first in the README (what this run is for)")
    gen.add_argument(
        "--store-label", default="", help="how the README names the store (default: the redacted URL)"
    )
    sc = sub.add_parser("score", help="score a filled workbook (or a directory of its CSVs)")
    sc.add_argument("path", type=pathlib.Path)
    sc.add_argument("--json", action="store_true")
    args = parser.parse_args(argv)

    if args.cmd == "score":
        score = score_rows(*read_filled(args.path))
        print(json.dumps(score.as_dict(), indent=1) if args.json else render_score(score))  # noqa: T201 -- CLI
        return 0

    if args.posture:
        os.environ["PLATFORM_POSTURE"] = args.posture
    from services.api import visibility
    from services.db.session import get_engine, get_sessionmaker
    from services.posture import platform_posture, publishable_reuse_classes
    from services.sites.switch import sites_enabled

    posture = platform_posture()
    if tuple(visibility.PUBLISHABLE_REUSE_CLASSES) != tuple(publishable_reuse_classes(posture)):
        parser.error("the visibility predicate was imported under another posture; run in a fresh process")
    if not sites_enabled():
        print("warning: SITES_ENABLED is off, so no site is served and the Sites sheet will be empty")  # noqa: T201
    out = args.out or pathlib.Path("data/eval/handcheck") / dt.datetime.now(dt.UTC).date().isoformat()
    store = args.store_label or _redacted(args.database_url)
    command = (
        f"python -m services.sites.handcheck generate --database-url <url> --out {out} --seed {args.seed}"
        f" --tier {args.tier} --posture {posture}"
    )
    engine = get_engine(args.database_url)
    session = get_sessionmaker(engine)()
    try:
        _read_only(session)
        sample = generate(
            session,
            out,
            seed=args.seed,
            entitlement=args.tier,
            n_sites=args.sites,
            n_merged=args.merged,
            largest=args.largest,
            posture=posture,
            store=store,
            base_url=args.base_url,
            command=command,
            note=args.note,
        )
    finally:
        session.rollback()
        session.close()
        engine.dispose()
    print(describe(sample))  # noqa: T201 -- CLI report
    print(f"written: {out / WORKBOOK}")  # noqa: T201
    return 0


if __name__ == "__main__":
    sys.exit(main())
